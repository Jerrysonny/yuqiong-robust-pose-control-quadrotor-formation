from __future__ import annotations

import csv
import hashlib
import json
import math
from copy import deepcopy
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
REPORT_ROOT = ROOT / "A8赛题核心工作区" / "仿真报告规划_20260712"
PACKAGE_ROOT = ROOT / "A8赛题核心工作区" / "A8比赛作品完整源文件_20260715"
TEMPLATE = (
    REPORT_ROOT
    / "08_第八章讨论结论与展望_20260714"
    / "A8仿真分析报告_第8章独立初稿_R2_20260714.docx"
)
PARAMETER_CSV = (
    PACKAGE_ROOT
    / "04_results"
    / "report_evidence"
    / "05_parameter_11_paired.csv"
)
STEP_CSV = (
    PACKAGE_ROOT
    / "04_results"
    / "report_evidence"
    / "02_standard_step_pid_vs_main.csv"
)
COMMON_CSV = (
    PACKAGE_ROOT
    / "04_results"
    / "report_evidence"
    / "03_five_controller_common_scenes.csv"
)
WORK = Path(__file__).resolve().parent
FIGURE = WORK / "figure_8_1_pid_base_main_table_compact.png"
FIGURE_PDF = WORK / "figure_8_1_pid_base_main_table_compact.pdf"
OUTPUT = WORK / "chapter8_candidate_03_compact_table.docx"
RECORD = WORK / "generation_record_compact_table.json"


def require_file(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)


def ensure_absent(path: Path) -> None:
    if path.exists():
        raise FileExistsError(path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def load_parameter_rows() -> list[dict[str, str]]:
    # 第8章只读取报告已采用的参数对照结果。
    with PARAMETER_CSV.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if len(rows) != 11:
        raise ValueError(f"Expected 11 parameter rows, found {len(rows)}")
    return rows


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def build_figure(step_rows: list[dict[str, str]], common_rows: list[dict[str, str]]) -> None:
    # 综合图优先展示官方PID与主算法的代表结果。
    for path in (FIGURE, FIGURE_PDF):
        ensure_absent(path)

    simsun_path = Path("C:/Windows/Fonts/simsun.ttc")
    require_file(simsun_path)
    font_manager.fontManager.addfont(str(simsun_path))
    simsun = font_manager.FontProperties(fname=str(simsun_path)).get_name()

    mpl.rcParams.update(
        {
            "font.family": simsun,
            "font.sans-serif": [simsun],
            "axes.unicode_minus": False,
            "font.size": 7.6,
            "axes.titlesize": 9.0,
            "axes.labelsize": 7.6,
            "xtick.labelsize": 6.8,
            "ytick.labelsize": 6.8,
            "axes.linewidth": 0.65,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )

    step_map = {
        (row["axis"], row["controller"]): float(row["overshoot_percent"])
        for row in step_rows
    }
    axes = ["X", "Y", "Z"]
    pid_overshoot = [step_map[(axis, "PID")] for axis in axes]
    main_overshoot = [step_map[(axis, "RA-GCA-CGHTE")] for axis in axes]

    common_map = {
        (row["case_id"], row["controller"]): float(row["tracking_rmse_m"])
        for row in common_rows
    }
    common_cases = ["Scene01", "Scene04", "Scene03", "Scene06b"]
    common_labels = ["分段爬升", "定点悬停", "8字轨迹", "三事件外扰"]
    pid_common = [common_map[(case, "PID")] for case in common_cases]
    base_common = [common_map[(case, "RA-GCA/Base")] for case in common_cases]
    main_common = [common_map[(case, "RA-GCA-CGHTE")] for case in common_cases]

    fig = plt.figure(figsize=(7.30, 4.15), constrained_layout=True, facecolor="white")
    outer = fig.add_gridspec(1, 3, width_ratios=[0.88, 1.30, 1.02], wspace=0.34)
    ax_a = fig.add_subplot(outer[0, 0])
    ax_b = fig.add_subplot(outer[0, 1])
    right = outer[0, 2].subgridspec(2, 1, height_ratios=[1.25, 0.75], hspace=0.16)
    ax_c = fig.add_subplot(right[0, 0])
    ax_c_note = fig.add_subplot(right[1, 0])

    blue = "#2F6B9A"
    teal = "#2F7D6D"
    amber = "#C58A2B"
    red = "#B6534C"
    green = "#3B7F5A"
    gray = "#7A7A7A"
    grid = "#D7D7D7"

    x_a = np.arange(len(axes))
    width_a = 0.34
    pid_bars = ax_a.bar(x_a - width_a / 2, pid_overshoot, width_a, color=blue, label="PID")
    main_bars = ax_a.bar(x_a + width_a / 2, main_overshoot, width_a, color=teal, label="RA-GCA-CGHTE")
    ax_a.set_xticks(x_a, axes)
    ax_a.set_ylim(0, 30)
    ax_a.set_ylabel("超调量（%）")
    ax_a.set_title("(a) 三轴阶跃超调", pad=7)
    ax_a.grid(axis="y", color=grid, linewidth=0.55, linestyle=(0, (2, 2)))
    ax_a.set_axisbelow(True)
    ax_a.legend(loc="upper center", bbox_to_anchor=(0.5, -0.09), ncol=1, frameon=False, fontsize=6.2)
    for bars in (pid_bars, main_bars):
        for bar in bars:
            value = bar.get_height()
            ax_a.text(
                bar.get_x() + bar.get_width() / 2,
                max(value + 0.55, 0.55),
                f"{value:.2f}" if value > 0 else "0",
                ha="center",
                va="bottom",
                fontsize=6.2,
            )

    y_b = np.arange(len(common_labels))
    height_b = 0.20
    ax_b.barh(y_b - height_b, pid_common, height_b, color=blue, label="PID")
    ax_b.barh(y_b, base_common, height_b, color=gray, label="RA-GCA/Base")
    main_bars_b = ax_b.barh(y_b + height_b, main_common, height_b, color=teal, label="RA-GCA-CGHTE")
    ax_b.set_yticks(y_b, common_labels)
    ax_b.invert_yaxis()
    ax_b.set_xlim(0, 0.38)
    ax_b.set_xticks([0, 0.1, 0.2, 0.3])
    ax_b.set_xlabel("位置RMSE（m）")
    ax_b.set_title("(b) 三种算法的共同任务", pad=7)
    ax_b.grid(axis="x", color=grid, linewidth=0.55, linestyle=(0, (2, 2)))
    ax_b.set_axisbelow(True)
    ax_b.legend(loc="upper center", bbox_to_anchor=(0.5, -0.09), ncol=1, frameon=False, fontsize=6.1)
    for bar, value in zip(main_bars_b, main_common):
        ax_b.text(value + 0.006, bar.get_y() + bar.get_height() / 2, f"{value:.3f}", va="center", ha="left", fontsize=5.9)

    distances = [0.3831758291349347, 0.6113274026131267]
    bars_c = ax_c.bar([0, 1], distances, width=0.56, color=[red, green], edgecolor="white", linewidth=0.5)
    ax_c.axhline(0.60, color="#333333", linewidth=0.9, linestyle=(0, (4, 2)))
    ax_c.set_xticks([0, 1], ["关闭", "开启"])
    ax_c.set_ylim(0, 0.72)
    ax_c.set_ylabel("最小间距（m）")
    ax_c.set_title("(c) PP-CBF安全收益与代价", pad=7)
    ax_c.grid(axis="y", color=grid, linewidth=0.55, linestyle=(0, (2, 2)))
    ax_c.set_axisbelow(True)
    ax_c.text(-0.40, 0.606, "0.60 m", ha="left", va="bottom", fontsize=6.4)
    for bar, value in zip(bars_c, distances):
        ax_c.text(bar.get_x() + bar.get_width() / 2, value + 0.014, f"{value:.4f}", ha="center", va="bottom", fontsize=6.8)

    ax_c_note.axis("off")
    ax_c_note.text(0.00, 0.98, "低于0.60 m累计时间", fontsize=6.3, color="#333333", ha="left", va="top")
    ax_c_note.text(0.00, 0.81, "2.04 s -> 0 s", fontsize=6.8, color=green, ha="left", va="top", fontweight="bold")
    ax_c_note.text(0.00, 0.63, "全局跟踪RMSE", fontsize=6.3, color="#333333", ha="left", va="top")
    ax_c_note.text(0.00, 0.46, "0.0672 -> 0.0833 m", fontsize=6.5, color=amber, ha="left", va="top")
    ax_c_note.text(0.00, 0.28, "相对队形RMSE", fontsize=6.3, color="#333333", ha="left", va="top")
    ax_c_note.text(0.00, 0.11, "0.0825 -> 0.2477 m", fontsize=6.5, color=amber, ha="left", va="top")

    for axis in (ax_a, ax_b, ax_c):
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)

    fig.savefig(FIGURE, dpi=420, bbox_inches="tight", facecolor="white")
    fig.savefig(FIGURE_PDF, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def set_style_fonts(style, east_asia: str = "宋体", latin: str = "Times New Roman") -> None:
    style.font.name = latin
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("ascii", latin), ("hAnsi", latin), ("eastAsia", east_asia), ("cs", latin)):
        rfonts.set(qn(f"w:{key}"), value)


def set_run_fonts(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"), ("eastAsia", "宋体"), ("cs", "Times New Roman")):
        rfonts.set(qn(f"w:{key}"), value)


def ensure_style(doc: Document, name: str, base: str = "Normal"):
    names = [style.name for style in doc.styles]
    if name in names:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    return style


def clear_body_keep_section(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for rel_id, relationship in list(doc.part.rels.items()):
        if relationship.reltype == RT.IMAGE:
            doc.part.drop_rel(rel_id)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._element):
        paragraph._element.remove(child)


def add_field(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])
    set_run_fonts(run, 9)


def set_paragraph_border(paragraph, side: str, color: str = "808080", size: str = "4") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), color)
    borders.append(edge)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style_fonts(normal)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 15, 0, 12),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 10, 4),
    ):
        style = doc.styles[name]
        set_style_fonts(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    set_style_fonts(caption)
    caption.font.size = Pt(10.5)
    caption.font.bold = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_with_next = True

    source = ensure_style(doc, "Source Note")
    set_style_fonts(source)
    source.font.size = Pt(9)
    source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    source.paragraph_format.first_line_indent = Pt(0)
    source.paragraph_format.left_indent = Cm(0.3)
    source.paragraph_format.right_indent = Cm(0.3)
    source.paragraph_format.space_after = Pt(6)
    source.paragraph_format.line_spacing = 1.0

    reference = ensure_style(doc, "Reference Entry")
    set_style_fonts(reference)
    reference.font.size = Pt(10.5)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.left_indent = Pt(24)
    reference.paragraph_format.first_line_indent = Pt(-24)
    reference.paragraph_format.space_after = Pt(3)
    reference.paragraph_format.line_spacing = 1.25
    reference.paragraph_format.widow_control = True
    reference.paragraph_format.keep_together = True

    conclusion = ensure_style(doc, "Conclusion Item")
    set_style_fonts(conclusion)
    conclusion.font.size = Pt(12)
    conclusion.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    conclusion.paragraph_format.left_indent = Pt(24)
    conclusion.paragraph_format.first_line_indent = Pt(-24)
    conclusion.paragraph_format.line_spacing = 1.5
    conclusion.paragraph_format.space_after = Pt(0)
    conclusion.paragraph_format.widow_control = True

    header = section.header
    hp = header.paragraphs[0]
    clear_paragraph(hp)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run("A8四旋翼无人机位姿控制系统仿真分析报告")
    set_run_fonts(run, 9)
    set_paragraph_border(hp, "bottom", color="B0B0B0", size="3")

    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("第 ")
    set_run_fonts(run, 9)
    add_field(fp, " PAGE ")
    run = fp.add_run(" 页")
    set_run_fonts(run, 9)


def add_heading(doc: Document, text: str, level: int):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_fonts(run, 15 if level == 1 else 14, True)
    return paragraph


def add_text(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_fonts(run, 12)


def add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Source Note")
    run = paragraph.add_run(text)
    set_run_fonts(run, 9)


def add_conclusion(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph(style="Conclusion Item")
    run = paragraph.add_run(f"（{number}）{text}")
    set_run_fonts(run, 12)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def set_cell_margins(cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_three_line_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, val, size in (
        ("top", "single", "12"),
        ("left", "nil", "0"),
        ("bottom", "single", "12"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tbl_pr.append(borders)

    header_borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")
    header_borders.append(bottom)
    for cell in table.rows[0].cells:
        tcpr = cell._tc.get_or_add_tcPr()
        old_cell = tcpr.find(qn("w:tcBorders"))
        if old_cell is not None:
            tcpr.remove(old_cell)
        tcpr.append(deepcopy(header_borders))


def add_table_81(doc: Document) -> None:
    caption = doc.add_paragraph(style="Caption")
    run = caption.add_run("表8-1 主要方法与代表结果")
    set_run_fonts(run, 10.5, True)

    headers = ["方法或比较", "代表结果"]
    rows = [
        [
            "RA-GCA-CGHTE / PID",
            "四类连续任务RMSE降低25.0%至91.2%；三轴阶跃超调明显减小",
        ],
        [
            "RA-GCA-CGHTE / RA-GCA/Base",
            "四个共同任务的RMSE均较低，降幅为2.3%至15.3%",
        ],
        [
            "RA-GCA-CGHTE参数变化",
            "11项设置的位置RMSE为0.01469 m至0.03368 m",
        ],
        [
            "PP-CBF开 / 关",
            "最小间距0.3832 m增至0.6113 m；低于0.60 m的时间2.04 s降至0 s；跟踪RMSE增加23.95%",
        ],
        [
            "三事件外扰",
            "CP-INDI为0.06212 m，RA-GCA-CGHTE为0.29490 m",
        ],
    ]

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(4.6), Cm(10.9)]
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(header)
        set_run_fonts(run, 9, True)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_data in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row_data)):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            set_run_fonts(run, 9)
        prevent_row_split(table.rows[-1])

    set_three_line_borders(table)
    add_note(doc, "注：结果仅对应本文相同任务和评价时间范围。")


def add_reference(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Reference Entry")
    run = paragraph.add_run(text)
    set_run_fonts(run, 10.5)


def build_document() -> None:
    # 第8章文档按固定章节结构、图表编号和样式生成。
    ensure_absent(OUTPUT)
    doc = Document(TEMPLATE)
    clear_body_keep_section(doc)
    setup_document(doc)

    add_heading(doc, "第8章 讨论、结论与展望", 1)
    add_text(
        doc,
        "本章综合单机跟踪、参数变化、外部扰动和三机编队结果，分析RA-GCA-CGHTE各组成部分的实际作用，并说明性能取舍与适用范围。讨论只使用前文已经给出的同条件结果，不以单一指标对五种控制方法作总体排名。",
    )

    add_heading(doc, "8.1 不同任务中的性能与方法作用", 2)
    add_text(
        doc,
        "RA-GCA-CGHTE由位置控制、约化姿态控制、约束感知分配和CGHTE推力尺度估计共同组成。位置控制根据期望合力确定推力方向，约化姿态控制优先对准与平移运动直接相关的机体推力轴，控制分配再将总推力及滚转、俯仰力矩转换为四路电机指令。该结构与几何控制、倾斜优先控制和约束控制分配的一般思路一致（Lee et al., 2010; Brescianini & D'Andrea, 2020; Johansen & Fossen, 2013）。CGHTE在满足稳定飞行条件时修正推力尺度，用于减小载荷或升力效能变化造成的高度误差。",
    )
    add_text(
        doc,
        "标准阶跃结果表明，PID在X、Y、Z三个方向的超调量分别为24.29%、24.28%和23.81%，RA-GCA-CGHTE对应为0、0和1.98%。由于PID阶跃数据没有与主算法一致的三维位置RMSE，本文只比较超调量，不计算阶跃RMSE改善率。在分段爬升、定点悬停、螺旋爬升和8字轨迹中，主算法相对PID的位置RMSE分别降低43.7%、91.2%、25.0%和70.0%。这些结果说明，推力方向控制与约化姿态调节能够兼顾静态保持、升降运动和连续曲线跟踪。",
    )
    add_text(
        doc,
        "图8-1将三类结果分开展示。左图比较PID与主算法的三轴阶跃超调，中图比较PID、RA-GCA/Base和主算法在四个共同任务中的位置RMSE，右图同时报告PP-CBF带来的安全收益和任务代价。三类指标的物理含义不同，因此分别分析。",
    )

    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = Pt(0)
    picture_paragraph.paragraph_format.space_before = Pt(3)
    picture_paragraph.paragraph_format.space_after = Pt(0)
    run = picture_paragraph.add_run()
    shape = run.add_picture(str(FIGURE), width=Cm(15.35))
    doc_pr = shape._inline.docPr
    doc_pr.set("name", "图8-1 不同任务中的主要结果与性能取舍")
    doc_pr.set("descr", "三部分综合图：PID与主算法的阶跃超调、PID和Base及主算法的共同任务结果、PP-CBF安全收益与跟踪代价")
    caption = doc.add_paragraph(style="Caption")
    run = caption.add_run("图8-1 不同任务中的主要结果与性能取舍")
    set_run_fonts(run, 10.5, True)
    add_note(
        doc,
        "注：(a)比较PID与主算法的三轴阶跃超调；(b)比较PID、RA-GCA/Base和主算法的四个共同任务；(c)中的0.60 m为本文采用的评价距离。",
    )

    add_text(
        doc,
        "四个共同任务提供了主算法与PID、RA-GCA/Base的统一比较。在分段爬升、定点悬停、8字轨迹和三事件外扰中，主算法相对RA-GCA/Base的位置RMSE分别降低4.7%、2.3%、6.1%和15.3%；相对PID分别降低43.7%、91.2%、70.0%和3.9%。这些结果只对应四个共同任务，不用于推断缺少同组数据的参数、风扰和传感器场景。",
    )
    add_text(
        doc,
        "主算法在11项参数设置中的位置RMSE为0.01469 m至0.03368 m，其中升力降低10%时为0.02251 m，载荷增加10%时为0.02152 m。CGHTE用于在稳定飞行条件下修正推力尺度，上述绝对结果表明集成后的主算法在受测载荷和升力效能变化下保持了较小的跟踪误差。由于PID和RA-GCA/Base没有同组11项参数结果，本章不计算参数场景的横向改善率。",
    )
    add_text(
        doc,
        "五种控制对象在共同任务中的结果显示，各方法的优势随任务变化。三事件外扰中，CP-INDI的位置RMSE为0.06212 m，低于RA-GCA-CGHTE的0.29490 m，说明增量动态逆在该突发扰动设置下具有明显优势（Smeur et al., 2016）。主算法在标准跟踪、参数变化和统一工程实现方面表现较均衡，但这一结果不能替代分场景比较。",
    )
    add_text(
        doc,
        "固定三机队形变换中，RA-GCA-CGHTE低层控制器取得0.0383 m的相对队形RMSE、0.0287 m的三机全局位置RMSE和0.8530 m的最小间距。近距离交汇时，开启PP-CBF使最小间距由0.3832 m增至0.6113 m，低于0.60 m的累计时间由2.04 s降为0 s。PP-CBF通过修正上层参考提高机间距离（Ames et al., 2017），同时使全局位置RMSE由0.0672 m增至0.0833 m，相对队形RMSE由0.0825 m增至0.2477 m。安全收益及其跟踪代价需要同时评价。",
    )

    add_table_81(doc)

    add_heading(doc, "8.2 性能取舍、控制能力与适用范围", 2)
    add_text(
        doc,
        "外扰结果反映了控制方法的任务差异。CP-INDI在三事件外扰中的局部优势值得保留。六组复合参数与外扰测试中，主算法的位置RMSE为0.210 m至0.317 m，恢复时间为1.61 s至1.72 s。系统能够在受测复合条件下恢复跟踪，但误差峰值和恢复速度仍有改进空间。",
    )
    add_text(
        doc,
        "在五个风扰初始相位下，主算法的位置RMSE均值为0.03216 m，结果范围为0.03196 m至0.03231 m；在含延迟与噪声的五个传感器设置下，均值为0.01866 m，结果范围为0.018657 m至0.018668 m。上述结果说明受测设置之间的波动较小。由于PID、RA-GCA/Base、CAP-ADRC和CP-INDI没有同组风扰、传感器及11项参数结果，本文不作这些场景的横向排名。",
    )
    add_text(
        doc,
        "RA-GCA-CGHTE控制位置、滚转和俯仰，持续偏航不属于本文验证范围。所用四旋翼模型未提供可持续调节的偏航反扭矩，因此相关结果不能解释为完整六自由度姿态跟踪。控制器包含约束感知分配和去饱和处理，但当前场景没有触发主动去饱和，其性能仍需在明确的执行器受限任务中测试。",
    )
    add_text(
        doc,
        "编队结论适用于固定三机、给定轨迹、同步状态输入和本文测试初值。0.60 m是本文采用的评价距离。通信延迟、丢包、更多成员、不同编队拓扑、硬件在环和实机飞行尚未开展，因此仿真结果主要说明方法在当前模型和测试条件下的可行性。",
    )

    add_heading(doc, "8.3 主要结论", 2)
    add_text(doc, "根据本文仿真结果，可以得到以下结论。")
    add_conclusion(
        doc,
        1,
        "RA-GCA-CGHTE在标准阶跃、分段爬升、定点悬停、螺旋爬升和8字轨迹中保持稳定跟踪。相对PID，四类连续任务的位置RMSE降低25.0%至91.2%，三轴阶跃超调也明显减小。",
    )
    add_conclusion(
        doc,
        2,
        "RA-GCA-CGHTE在11项受测参数设置中的位置RMSE为0.01469 m至0.03368 m，升力降低10%和载荷增加10%时分别为0.02251 m和0.02152 m。这些结果支持主算法在本文参数变化范围内的跟踪能力，不构成与缺少同组数据算法的横向排名。",
    )
    add_conclusion(
        doc,
        3,
        "RA-GCA-CGHTE能够支持固定三机队形变换。PP-CBF在相同低层控制器下提高了近距离交汇的最小间距，并消除了低于本文评价距离的累计时间；代价是全局位置和相对队形误差增加。",
    )
    add_conclusion(
        doc,
        4,
        "不同控制方法存在任务相关的优势。CP-INDI在受测三事件外扰中优于主算法，主算法在复合测试中的恢复时间为1.61 s至1.72 s。本文结论限于位置与约化姿态控制、给定仿真模型和测试场景，不延伸到持续偏航、未触发的主动去饱和或实机条件。",
    )

    add_heading(doc, "8.4 后续工作与未来展望", 2)
    add_text(
        doc,
        "后续首先开展硬件在环和实机飞行测试，测量控制周期、计算负载、传感器误差和执行器差异对结果的影响，并使用多次重复实验给出均值、离散程度和置信区间。参数变化与外扰同时出现时，应进一步降低误差峰值并缩短恢复时间。",
    )
    add_text(
        doc,
        "其次，在包含桨叶反扭矩的模型和实物平台上补充偏航控制；设计能够稳定触发电机上下限的任务，评价约束分配与主动去饱和对姿态保持和电机使用的实际作用。还可结合外扰观测与推力尺度估计，提高突发扰动和参数变化同时出现时的控制性能。",
    )
    add_text(
        doc,
        "多机部分将扩展到更多成员、不同拓扑以及通信延迟和丢包条件，并研究安全距离、跟踪精度和控制修正之间的自适应调节。通过仿真、硬件在环和实机逐级验证，可进一步评估该方法在巡检、运输和多机协同任务中的应用条件。",
    )

    reference_heading = add_heading(doc, "参考文献", 1)
    reference_heading.paragraph_format.page_break_before = False
    add_reference(
        doc,
        "Ames, A. D., Xu, X., Grizzle, J. W., & Tabuada, P. (2017). Control barrier function based quadratic programs for safety-critical systems. IEEE Transactions on Automatic Control, 62(8), 3861-3876. https://doi.org/10.1109/TAC.2016.2638961",
    )
    add_reference(
        doc,
        "Brescianini, D., & D'Andrea, R. (2020). Tilt-prioritized quadrocopter attitude control. IEEE Transactions on Control Systems Technology, 28(2), 376-387. https://doi.org/10.1109/TCST.2018.2873224",
    )
    add_reference(
        doc,
        "Johansen, T. A., & Fossen, T. I. (2013). Control allocation - A survey. Automatica, 49(5), 1087-1103. https://doi.org/10.1016/j.automatica.2013.01.035",
    )
    add_reference(
        doc,
        "Lee, T., Leok, M., & McClamroch, N. H. (2010). Geometric tracking control of a quadrotor UAV on SE(3). 49th IEEE Conference on Decision and Control, 5420-5425. https://doi.org/10.1109/CDC.2010.5717652",
    )
    add_reference(
        doc,
        "Smeur, E. J. J., Chu, Q., & de Croon, G. C. H. E. (2016). Adaptive incremental nonlinear dynamic inversion for attitude control of micro air vehicles. Journal of Guidance, Control, and Dynamics, 39(3), 450-461. https://doi.org/10.2514/1.G001490",
    )

    core = doc.core_properties
    core.title = "A8仿真分析报告第8章：讨论、结论与展望"
    core.subject = "RA-GCA-CGHTE仿真结果讨论、主要结论与后续工作"
    core.author = "A8项目组"
    core.keywords = "四旋翼; RA-GCA-CGHTE; PP-CBF; 仿真分析; 讨论与结论"
    core.comments = "基于第1至第7章最新稿重写的独立第8章"

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_run_fonts(run, run.font.size.pt if run.font.size else None, run.bold)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_fonts(run, run.font.size.pt if run.font.size else 9, run.bold)

    doc.save(OUTPUT)


def validate() -> dict[str, object]:
    # 验收同时检查内容口径、图表数量和基础排版属性。
    doc = Document(OUTPUT)
    paragraph_text = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    table_text = [cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells if cell.text.strip()]
    full_text = "\n".join(paragraph_text + table_text)
    errors: list[str] = []

    required_headings = [
        "第8章 讨论、结论与展望",
        "8.1 不同任务中的性能与方法作用",
        "8.2 性能取舍、控制能力与适用范围",
        "8.3 主要结论",
        "8.4 后续工作与未来展望",
        "参考文献",
    ]
    for heading in required_headings:
        if paragraph_text.count(heading) != 1:
            errors.append(f"Heading count is not 1: {heading}")

    required_phrases = [
        "RA-GCA-CGHTE",
        "43.7%",
        "91.2%",
        "25.0%",
        "70.0%",
        "0.06212 m",
        "0.29490 m",
        "0.3832 m",
        "0.6113 m",
        "0.0672 m增至0.0833 m",
        "0.0825 m增至0.2477 m",
        "0.01469 m至0.03368 m",
        "持续偏航不属于本文验证范围",
    ]
    for phrase in required_phrases:
        if phrase not in full_text:
            errors.append(f"Required phrase missing: {phrase}")

    forbidden = [
        "RA-GCA + CG-HTE",
        "RA-GCA + CG-HTE 1.0.0",
        "V9",
        "HTE-F0",
        "candidate",
        "SHA256",
        "交接",
        "冻结",
        "内部Score",
        "硬门",
        "证据链",
        "反馈链",
        "控制门",
        "证据覆盖协议",
        "V30",
        "H60",
        "Wilson",
        "bootstrap",
        "0.432%",
    ]
    for phrase in forbidden:
        if phrase in full_text:
            errors.append(f"Forbidden phrase present: {phrase}")

    if "不是" in full_text and "而是" in full_text:
        errors.append("Potential avoidable 不是...而是 construction")
    if "不再" in full_text and "而是" in full_text:
        errors.append("Potential avoidable 不再...而是 construction")
    if len(doc.inline_shapes) != 1:
        errors.append(f"Expected 1 figure, found {len(doc.inline_shapes)}")
    if len(doc.tables) != 1:
        errors.append(f"Expected 1 table, found {len(doc.tables)}")
    if doc.tables and (len(doc.tables[0].rows) != 6 or len(doc.tables[0].columns) != 2):
        errors.append("Table 8-1 dimensions are not 6x2")

    section = doc.sections[0]
    expected = {
        "page_width": 21.0,
        "page_height": 29.7,
        "left_margin": 3.0,
        "right_margin": 2.5,
        "top_margin": 2.5,
        "bottom_margin": 2.5,
    }
    for attribute, expected_cm in expected.items():
        actual_cm = getattr(section, attribute).cm
        if abs(actual_cm - expected_cm) > 0.02:
            errors.append(f"Unexpected {attribute}: {actual_cm:.3f} cm")

    if OUTPUT.stat().st_size < 200_000:
        errors.append("DOCX is unexpectedly small")
    if FIGURE.stat().st_size < 150_000:
        errors.append("Figure is unexpectedly small")

    result: dict[str, object] = {
        "output": str(OUTPUT),
        "output_bytes": OUTPUT.stat().st_size,
        "output_sha256": sha256(OUTPUT),
        "figure": str(FIGURE),
        "figure_bytes": FIGURE.stat().st_size,
        "figure_sha256": sha256(FIGURE),
        "template_sha256": sha256(TEMPLATE),
        "parameter_csv_sha256": sha256(PARAMETER_CSV),
        "step_csv_sha256": sha256(STEP_CSV),
        "common_csv_sha256": sha256(COMMON_CSV),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "figures": len(doc.inline_shapes),
        "errors": errors,
    }
    RECORD.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if errors:
        raise RuntimeError("; ".join(errors))
    return result


def main() -> None:
    for path in (TEMPLATE, PARAMETER_CSV, STEP_CSV, COMMON_CSV):
        require_file(path)
    for path in (FIGURE, FIGURE_PDF, OUTPUT, RECORD):
        ensure_absent(path)
    load_parameter_rows()
    step_rows = load_csv(STEP_CSV)
    common_rows = load_csv(COMMON_CSV)
    build_figure(step_rows, common_rows)
    build_document()
    result = validate()
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

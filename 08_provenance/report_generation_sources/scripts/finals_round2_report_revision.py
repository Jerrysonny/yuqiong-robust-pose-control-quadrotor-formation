from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


REPORT_NAME = "仿真报告_驭穹稳控_基于MWORKS的四旋翼鲁棒位姿控制与编队安全仿真.docx"
EXPECTED_SOURCE_SHA256 = "B53B109B2F09E6C23A2A9A9415B8CE79F75A9CB030D2ED9A5A6A887A3DA18642"
FIGURE_WIDTH = Cm(16.2)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def find_starts(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def find_caption(document: Document, prefix: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption" and paragraph.text.strip().startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def preceding_drawing_paragraph(caption: Paragraph) -> Paragraph:
    node = caption._p.getprevious()
    checked = 0
    while node is not None and checked < 12:
        if node.tag == qn("w:p") and node.xpath(".//w:drawing"):
            return Paragraph(node, caption._parent)
        if node.tag == qn("w:p"):
            text = "".join(node.itertext()).strip()
            if text.startswith(("图", "表")):
                break
        node = node.getprevious()
        checked += 1
    raise RuntimeError(f"No drawing paragraph found before caption: {caption.text}")


def replace_picture_before_caption(document: Document, caption_prefix: str, image_path: Path) -> None:
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    caption = find_caption(document, caption_prefix)
    image_paragraph = preceding_drawing_paragraph(caption)
    clear_paragraph(image_paragraph)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(image_path), width=FIGURE_WIDTH)
    caption.paragraph_format.keep_with_next = True


def set_cell_margins(cell, *, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_border(container, side: str, *, value: str, size: int = 0, color: str = "000000") -> None:
    border_tag = "w:tblBorders" if container.tag == qn("w:tblPr") else "w:tcBorders"
    borders = container.find(qn(border_tag))
    if borders is None:
        borders = OxmlElement(border_tag)
        container.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), value)
    if value != "nil":
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_white_shading(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), "FFFFFF")
    shading.set(qn("w:val"), "clear")


def remove_fixed_row_height(row) -> None:
    properties = row._tr.get_or_add_trPr()
    for height in list(properties.findall(qn("w:trHeight"))):
        properties.remove(height)


def style_three_line_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    properties = table._tbl.tblPr
    for side in ("left", "right", "insideH", "insideV"):
        set_border(properties, side, value="nil")
    set_border(properties, "top", value="single", size=12)
    set_border(properties, "bottom", value="single", size=12)

    for row_index, row in enumerate(table.rows):
        remove_fixed_row_height(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_white_shading(cell)
            cell_properties = cell._tc.get_or_add_tcPr()
            for side in ("left", "right", "top", "insideH", "insideV"):
                set_border(cell_properties, side, value="nil")
            set_border(cell_properties, "bottom", value="single" if row_index == 0 else "nil", size=6)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0


def numbered_tables(document: Document):
    tables = [table for table in document.tables if len(table.rows) > 1]
    if len(tables) != 28:
        raise RuntimeError(f"Expected 28 numbered tables, found {len(tables)}")
    return tables


def find_table_by_header(document: Document, header: list[str]):
    matches = []
    for table in document.tables:
        if not table.rows:
            continue
        values = [" ".join(cell.text.split()) for cell in table.rows[0].cells]
        if values == header:
            matches.append(table)
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table with header {header!r}, found {len(matches)}")
    return matches[0]


def rewrite_section_6_3(document: Document) -> None:
    set_text(
        find_starts(document, "ARDG补偿仅开放三个物理含义明确的参数"),
        "ARDG补偿只优化三个物理含义明确的参数：补偿增益Kc、补偿幅值上限Cmax和单步变化上限Rmax。参数整定采用受约束多目标贝叶斯优化：先用最大最小拉丁超立方在允许范围内均匀布置初始候选，再用高斯过程和RBF代理模型根据已有仿真结果预测新候选表现，并以约束期望超体积改进准则兼顾性能提升与搜索覆盖；最后在非支配候选附近进行小范围局部细化。",
    )
    set_text(
        find_starts(document, "最终定型参数为Kc"),
        "全局阶段评估14组候选，局部阶段评估8组候选，两个核心角扰动场景共完成44次筛选仿真；最终参数随后在两个场景中再次确认。最终参数为Kc = 0.89894、Cmax = 3.2953×10⁻⁴ N·m、Rmax = 4.1958×10⁻⁵ N·m/采样。优化同时考察姿态事件IAE、四电机总变差、最坏场景退化和安全约束，不使用加权总分，也不宣称全局最优。",
    )
    set_text(find_caption(document, "图6-5"), "图6-5  参数变化下的预设工况与随机配对结果")
    set_text(
        find_starts(document, "注：圆点为14组联合搜索候选"),
        "注：(a)为ARDG-RGPC在11项预设参数工况中的位置RMSE；(b)比较20组随机参数下官方PID与ARDG-RGPC的RMSE分布；(c)、(d)分别给出ARDG-RGPC相对官方PID的位置RMSE和峰值位置误差逐组降幅。",
    )
    set_text(
        find_starts(document, "注：参数定型使用两个核心角扰动场景"),
        "注：参数优化使用两个核心角扰动场景；最终参数确认后再进入代表性回归和工程验证。",
    )
    set_text(find_caption(document, "图6-6"), "图6-6  ARDG-RGPC参数优化过程与最终定型")
    set_text(
        find_starts(document, "注：(a)给出11项预设参数相对RA-GCA/Base的逐工况结果"),
        "注：(a)展示局部候选在三个调节参数空间中的分布；(b)、(c)分别给出姿态事件IAE和四电机总变差的候选筛选结果。星号表示最终定型参数。",
    )

    table = find_table_by_header(document, ["优化阶段", "候选组数", "核心场景仿真", "说明"])
    rows = [
        ["初始与全局搜索", "14", "28", "最大最小拉丁超立方布点，并用受约束多目标贝叶斯优化选择候选"],
        ["局部细化", "8", "16", "在非支配候选附近缩小范围继续搜索"],
        ["最终参数确认", "1", "2", "复核两个核心场景，随后进入代表性回归和工程验证"],
    ]
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value


def rewrite_public_diagnostics(document: Document) -> None:
    replacements = {
        "ARDG-RGPC采用11维参考、18维状态、2维分配上下限、4维转子转速、1维使能和1维时间输入":
            "ARDG-RGPC采用11维参考、18维状态、2维分配上下限、4维转子转速、1维使能和1维时间输入，输出4维电机命令与16维公共诊断。角残差的激活状态、融合系数、补偿力矩、残差力矩、通道模式和安全状态另作为结果记录，便于解释补偿何时介入。控制周期为0.01 s。Sysplorer建立四旋翼机体、电机和传感器模型，Syslab读取仿真结果并计算评价指标。",
        "滚转与俯仰补偿状态通过独立诊断CSV记录":
            "滚转与俯仰补偿状态作为单独的结果记录，公共16维诊断输出保持不变。记录内容包括激活状态、融合系数、两轴补偿力矩、两轴残差力矩、补偿通道模式和安全有效状态；独立偏航力矩不在本文控制范围内。",
        "正式Sysblock控制器保留11维参考":
            "正式Sysblock控制器保留11维参考、18维状态和2维分配上下限输入，并增加4维转子转速、1维使能和1维时间输入。公共16维诊断输出保持不变，角残差状态另作结果记录，因此不会改变原有公共接口。",
        "独立诊断CSV向用户公开":
            "角残差状态记录包括激活状态、融合系数、滚转/俯仰补偿力矩、滚转/俯仰残差力矩、补偿通道模式和安全有效状态，用于判断补偿是否介入以及安全旁路是否触发，不参与四旋翼动力学计算。",
        "表C-2按模型端口顺序列出":
            "表C-2按模型端口顺序列出ARDG-RGPC与PP-CBF的输入、控制输出和诊断信息。公共诊断保持16维，角残差状态单独列出；这些诊断量只用于结果解释和安全检查，不参与四旋翼动力学计算。",
        "两个机体系角扰动场景均记录独立诊断CSV":
            "两个机体系角扰动场景均记录角残差状态。图C-4展示角残差、受限补偿、守卫激活和安全有效状态的对应关系；安全有效状态为1表示该采样点通过数值、力矩、电机分配和旁路检查。",
    }
    for prefix, text in replacements.items():
        set_text(find_starts(document, prefix), text)

    set_text(
        find_starts(document, "注：示意依据正式Sysblock源码"),
        "注：示意依据正式Sysblock模型和公开接口绘制，图中采用功能名称说明各模块作用。",
    )

    table = find_table_by_header(document, ["模块", "输入", "控制输出", "诊断信息"])
    table.cell(2, 0).text = "角残差状态记录"
    table.cell(2, 1).text = "控制器内部计算状态"
    table.cell(2, 2).text = "不参与控制输出"
    table.cell(2, 3).text = "激活、融合、两轴补偿、两轴残差、通道模式和安全有效状态"

    interface_table = find_table_by_header(document, ["项目", "设定值", "说明"])
    for row in interface_table.rows:
        if row.cells[0].text.strip() == "接口":
            row.cells[2].text = "角残差状态另作结果记录，不改变公共诊断输出"
        if row.cells[0].text.strip() == "ARDG补偿":
            row.cells[2].text = "通过受约束多目标贝叶斯优化确定"

    port_table = find_table_by_header(document, ["接口", "维数", "分量顺序", "用途"])
    for row in port_table.rows:
        if row.cells[0].text.strip() == "诊断":
            row.cells[1].text = "公共16维 + 角残差状态记录"


def rewrite_other_internal_wording(document: Document) -> None:
    set_text(
        find_starts(document, "联合搜索使用14组候选"),
        "参数优化评估14组全局候选和8组局部候选，两个核心场景共完成44次筛选仿真。最终参数在两个核心场景中再次确认后，才用于三事件外力扰动和18项代表性回归；该过程不宣称全局最优。",
    )
    set_text(
        find_starts(document, "双精度生成代码共形成11项冻结工件"),
        "双精度代码生成共形成11项可复核输出。生成的控制器Step(void)函数经10000次预热、1000000次正式计时，P99为0.0002 ms、最大值为0.1476 ms，测量过程动态分配为0；电机剪裁和不可行分配旁路通过静态集成检查。",
    )


def fix_local_pagination(document: Document) -> None:
    # The source caption forces a new page and leaves the preceding result paragraph isolated.
    find_starts(document, "表6-4").paragraph_format.page_break_before = False


def replace_figures(document: Document, figure_dir: Path) -> None:
    for caption, filename in (
        ("图3-1", "FIG_3-1.png"),
        ("图4-1", "FIG_4-1.png"),
        ("图6-5", "FIG_6-5.png"),
        ("图6-6", "FIG_6-6.png"),
        ("图C-1", "FIG_C-1.png"),
        ("图C-2", "FIG_C-2.png"),
    ):
        replace_picture_before_caption(document, caption, figure_dir / filename)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    for node in settings.findall(qn("w:updateFields")):
        settings.remove(node)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def validate(document: Document) -> None:
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        full_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    forbidden = [
        "ardg_diagnostics.csv",
        "独立诊断CSV",
        "独立CSV",
        "ardg_safety_valid",
        "冻结后复核",
        "冻结后进入",
        "冻结工件",
        "圆点为14组联合搜索候选",
    ]
    found = [token for token in forbidden if token in full_text]
    if found:
        raise RuntimeError(f"Internal-only wording remains: {found}")
    if "图6-5  参数变化下的预设工况与随机配对结果" not in full_text:
        raise RuntimeError("Figure 6-5 caption was not updated")
    if "图6-6  ARDG-RGPC参数优化过程与最终定型" not in full_text:
        raise RuntimeError("Figure 6-6 caption was not updated")
    if "受约束多目标贝叶斯优化" not in full_text:
        raise RuntimeError("Concrete optimization method is missing")
    if len(document.tables) != 35 or len(document.inline_shapes) != 31:
        raise RuntimeError(
            f"Document object count changed: tables={len(document.tables)}, inline_shapes={len(document.inline_shapes)}"
        )
    for table in numbered_tables(document):
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            raise RuntimeError("A numbered table is not centered")
        for row in table.rows:
            for cell in row.cells:
                if cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER:
                    raise RuntimeError("A numbered table cell is not vertically centered")
                for paragraph in cell.paragraphs:
                    if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                        raise RuntimeError("A numbered table cell is not horizontally centered")


def build(root: Path, output: Path) -> None:
    source = root / REPORT_NAME
    figure_dir = root / "05_visuals/assets/report_final_20260816"
    if sha256(source) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("Source report hash mismatch; refusing incremental edit")
    document = Document(source)
    rewrite_section_6_3(document)
    rewrite_public_diagnostics(document)
    rewrite_other_internal_wording(document)
    fix_local_pagination(document)
    replace_figures(document, figure_dir)
    for table in numbered_tables(document):
        style_three_line_table(table)
    set_update_fields(document)
    validate(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError(f"refusing to overwrite candidate: {output}")
    document.save(output)
    print(f"source_sha256={sha256(source)}")
    print(f"output={output}")
    print(f"output_sha256={sha256(output)}")


def check(output: Path) -> None:
    document = Document(output)
    validate(document)
    print(f"candidate_check=pass")
    print(f"candidate_sha256={sha256(output)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    root = args.root.resolve()
    output = args.output.resolve()
    if args.check:
        check(output)
    else:
        build(root, output)


if __name__ == "__main__":
    main()

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from finals_round2_report_revision import (
    find_caption,
    find_starts,
    numbered_tables,
    replace_picture_before_caption,
    remove_fixed_row_height,
    set_border,
    set_cell_margins,
    set_white_shading,
    set_text,
    set_update_fields,
)


REPORT_NAME = "仿真报告_驭穹稳控_基于MWORKS的四旋翼鲁棒位姿控制与编队安全仿真.docx"
EXPECTED_SOURCE_SHA256 = "B53B109B2F09E6C23A2A9A9415B8CE79F75A9CB030D2ED9A5A6A887A3DA18642"
PUBLIC_BASELINE = "几何控制基线"
GITHUB_URL = "https://github.com/Jerrysonny/yuqiong-robust-pose-control-quadrotor-formation。"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def style_three_line_table(table) -> None:
    """Apply visible top, header-separator, and bottom rules at cell level."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_properties = table._tbl.tblPr
    for side in ("left", "right", "insideH", "insideV"):
        set_border(table_properties, side, value="nil")
    set_border(table_properties, "top", value="nil")
    set_border(table_properties, "bottom", value="nil")

    last_row = len(table.rows) - 1
    for row_index, row in enumerate(table.rows):
        remove_fixed_row_height(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_white_shading(cell)
            cell_properties = cell._tc.get_or_add_tcPr()
            for side in ("left", "right", "top", "bottom", "insideH", "insideV"):
                set_border(cell_properties, side, value="nil")
            if row_index == 0:
                set_border(cell_properties, "top", value="single", size=12)
                set_border(cell_properties, "bottom", value="single", size=6)
            if row_index == last_row:
                set_border(cell_properties, "bottom", value="single", size=12)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0


def replace_in_runs(paragraph, old: str, new: str) -> int:
    replacements = 0
    for run in paragraph.runs:
        if old in run.text:
            replacements += run.text.count(old)
            run.text = run.text.replace(old, new)
    return replacements


def replace_public_baseline_name(document: Document) -> int:
    replacements = 0
    for paragraph in document.paragraphs:
        replacements += replace_in_runs(paragraph, "RA-GCA/Base", PUBLIC_BASELINE)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replacements += replace_in_runs(paragraph, "RA-GCA/Base", PUBLIC_BASELINE)
    if replacements != 19:
        raise RuntimeError(f"Expected 19 public baseline replacements, found {replacements}")
    return replacements


def rewrite_parameter_optimization(document: Document) -> None:
    set_text(
        find_starts(document, "ARDG补偿只优化三个物理含义明确的参数"),
        "ARDG补偿只整定三个物理含义明确的参数：补偿增益Kc、补偿幅值上限Cmax和单步变化上限Rmax。参数整定采用受约束多目标贝叶斯优化：先用拉丁超立方覆盖允许范围，再根据已有仿真结果选择更有希望的候选，最后在综合表现较好的区域进行局部细化。优化同时考虑姿态误差、电机指令平滑性和安全约束。",
    )
    set_text(
        find_starts(document, "全局阶段评估14组候选"),
        "全局阶段评估14组候选，局部阶段评估8组候选，两个角扰动场景共完成44次筛选仿真；最终参数随后在两个场景中再次确认。最终取Kc = 0.89894、Cmax = 3.2953×10⁻⁴ N·m、Rmax = 4.1958×10⁻⁵ N·m/采样。该结果是在给定范围和约束下得到的工程定型参数，不表述为无边界的全局最优。",
    )
    set_text(find_caption(document, "图6-6"), "图6-6  ARDG-RGPC参数优化过程与最终定型")
    set_text(
        find_starts(document, "注：(a)给出11项预设参数相对"),
        "注：(a)展示8组局部候选在三个调节参数空间中的分布；(b)、(c)分别给出姿态事件IAE和四电机总变差的筛选结果。虚线表示局部搜索起点，星号表示最终定型参数。",
    )

    table = next(
        table
        for table in document.tables
        if table.rows
        and [" ".join(cell.text.split()) for cell in table.rows[0].cells]
        == ["优化阶段", "候选组数", "核心场景仿真", "说明"]
    )
    table.cell(1, 3).text = "拉丁超立方覆盖参数范围，并用受约束多目标贝叶斯优化选择候选"
    table.cell(2, 3).text = "在综合表现较好的候选附近缩小搜索范围"
    table.cell(3, 3).text = "在两个角扰动场景复核后，再进入代表性回归和工程验证"


def remove_redundant_code_availability(document: Document) -> None:
    """Keep the Chapter 1 repository citation and remove the repeated Chapter 8 copy."""
    label = "代码可获得性：本文使用的团队自研代码与复现材料可从GitHub仓库获取："
    label_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == label]
    url_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == GITHUB_URL]
    if len(label_paragraphs) != 1 or len(url_paragraphs) != 2:
        raise RuntimeError(
            "Unexpected code-availability structure: "
            f"labels={len(label_paragraphs)}, repository_urls={len(url_paragraphs)}"
        )
    label_paragraph = label_paragraphs[0]
    repeated_url = url_paragraphs[-1]
    if label_paragraph._p.getnext() is not repeated_url._p:
        raise RuntimeError("Repeated repository URL is not adjacent to the Chapter 8 label")
    label_paragraph._p.getparent().remove(label_paragraph._p)
    repeated_url._p.getparent().remove(repeated_url._p)


def replace_figures(document: Document, figure_dir: Path) -> None:
    for caption, filename in (
        ("图3-1", "FIG_3-1.png"),
        ("图4-1", "FIG_4-1.png"),
        ("图6-4", "FIG_6-4.png"),
        ("图6-5", "FIG_6-5.png"),
        ("图6-6", "FIG_6-6.png"),
        ("图8-1", "FIG_8-1.png"),
        ("图C-1", "FIG_C-1.png"),
        ("图C-2", "FIG_C-2.png"),
    ):
        replace_picture_before_caption(document, caption, figure_dir / filename)


def full_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def cell_border_value(cell, side: str) -> str | None:
    borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    if borders is None:
        return None
    border = borders.find(qn(f"w:{side}"))
    return None if border is None else border.get(qn("w:val"))


def validate(document: Document) -> None:
    text = full_text(document)
    forbidden = [
        "RA-GCA/Base",
        "RA-GCA-CGHTE",
        "v914",
        "v936",
        "L06",
        "J03",
        "97406",
        "ardg_diagnostics.csv",
        "独立诊断CSV",
        "冻结后复核",
        "冻结工件",
    ]
    found = [token for token in forbidden if token in text]
    if found:
        raise RuntimeError(f"Internal or obsolete wording remains: {found}")
    required = [
        "几何控制基线",
        "受约束多目标贝叶斯优化",
        "图6-5  参数变化下的预设工况与随机配对结果",
        "图6-6  ARDG-RGPC参数优化过程与最终定型",
        "20胜、0负、0平",
        "中位三维位置RMSE降幅为87.6345%",
        GITHUB_URL,
    ]
    missing = [token for token in required if token not in text]
    if missing:
        raise RuntimeError(f"Required report evidence is missing: {missing}")
    if len(document.tables) != 35 or len(document.inline_shapes) != 31:
        raise RuntimeError(
            f"Document object count changed: tables={len(document.tables)}, inline_shapes={len(document.inline_shapes)}"
        )
    if text.count(GITHUB_URL) != 1:
        raise RuntimeError(f"Expected one public repository URL, found {text.count(GITHUB_URL)}")
    tables = numbered_tables(document)
    if len(tables) != 28:
        raise RuntimeError(f"Expected 28 numbered tables, found {len(tables)}")
    for table in tables:
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            raise RuntimeError("A numbered table is not centered")
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                if cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER:
                    raise RuntimeError("A numbered table cell is not vertically centered")
                if row_index == 0:
                    if cell_border_value(cell, "top") != "single":
                        raise RuntimeError("A numbered table is missing its top rule")
                    if cell_border_value(cell, "bottom") != "single":
                        raise RuntimeError("A numbered table is missing its header separator")
                if row_index == len(table.rows) - 1 and cell_border_value(cell, "bottom") != "single":
                    raise RuntimeError("A numbered table is missing its bottom rule")
                for paragraph in cell.paragraphs:
                    if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                        raise RuntimeError("A numbered table cell is not horizontally centered")


def build(root: Path, figure_dir: Path, output: Path) -> None:
    source = root / REPORT_NAME
    if sha256(source) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("Source report hash mismatch; refusing incremental edit")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite candidate: {output}")
    required_figures = [
        "FIG_3-1.png",
        "FIG_4-1.png",
        "FIG_6-4.png",
        "FIG_6-5.png",
        "FIG_6-6.png",
        "FIG_8-1.png",
        "FIG_C-1.png",
        "FIG_C-2.png",
    ]
    missing_figures = [name for name in required_figures if not (figure_dir / name).is_file()]
    if missing_figures:
        raise FileNotFoundError(f"Candidate figures are missing: {missing_figures}")

    document = Document(source)
    replace_public_baseline_name(document)
    rewrite_parameter_optimization(document)
    remove_redundant_code_availability(document)
    replace_figures(document, figure_dir)
    for table in numbered_tables(document):
        style_three_line_table(table)
    set_update_fields(document)
    validate(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"source_sha256={sha256(source)}")
    print(f"candidate_sha256={sha256(output)}")
    print(f"output={output}")


def check(output: Path) -> None:
    document = Document(output)
    validate(document)
    print("candidate_check=pass")
    print(f"candidate_sha256={sha256(output)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path)
    parser.add_argument("--figure-dir", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    output = args.output.resolve()
    if args.check:
        check(output)
        return
    if args.root is None or args.figure_dir is None:
        parser.error("--root and --figure-dir are required unless --check is used")
    build(args.root.resolve(), args.figure_dir.resolve(), output)


if __name__ == "__main__":
    main()

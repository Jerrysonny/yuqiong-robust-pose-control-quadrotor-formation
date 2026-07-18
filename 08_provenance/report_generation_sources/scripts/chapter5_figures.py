from __future__ import annotations

import ast
import csv
import hashlib
import json
import math
import os
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from lxml import etree
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\Users\admin\Desktop\A8-四旋翼无人机位姿控制系统设计优化")
PLAN = ROOT / "A8赛题核心工作区" / "仿真报告规划_20260712"
BASELINE = PLAN / "A8仿真分析报告高质量精简大纲_最终算法收口版_20260714.docx"
WORK = PLAN / "05_第五章冻结内容" / "20260715_读者表达修订"
OUTPUT_DIR = WORK
OUTPUT = OUTPUT_DIR / "A8仿真分析报告_第5章读者表达修订稿_20260715.docx"
BUILD = WORK / "03_assets"
EVIDENCE_OUTPUT = WORK / "02_syslab_evidence"
EVIDENCE_INPUT = EVIDENCE_OUTPUT / "source_inputs"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
FINAL_PACKAGE = ROOT / "A8赛题核心工作区" / "A8比赛作品完整源文件_20260715"
RAW = FINAL_PACKAGE / "06_supplementary_evidence" / "ra_gca_cg_hte_32_raw"
STEP_X = RAW / "Scene01S_X.csv"
STEP_Y = RAW / "Scene01S_Y.csv"
STEP_Z = RAW / "Scene01S_Z.csv"
SCENE_01 = RAW / "Scene01.csv"
SCENE_02 = RAW / "Scene02.csv"
SCENE_03 = RAW / "Scene03.csv"
SCENE_06B = RAW / "Scene06b.csv"
PARAMETER_CONFIG = FINAL_PACKAGE / "02_scripts" / "sysplorer" / "run_main_controller.py"
PACKAGE_MANIFEST = FINAL_PACKAGE / "SHA256_MANIFEST.json"
COVERAGE_SOURCE = EVIDENCE_INPUT / "01_five_controller_scene_coverage.csv"
STEP_SOURCE = EVIDENCE_INPUT / "02_standard_step_pid_vs_main.csv"
COMMON_SOURCE = EVIDENCE_INPUT / "03_five_controller_common_scenes.csv"
EVOLUTION_SOURCE = EVIDENCE_INPUT / "04_algorithm_evolution_ablation.csv"
PARAMETER_SOURCE = EVIDENCE_INPUT / "05_parameter_11_paired.csv"
REGRESSION_SOURCE = EVIDENCE_INPUT / "06_project_regression_32_summary.json"
FORMATION_SOURCE = EVIDENCE_INPUT / "09_formation_pp_cbf.csv"
CLAIM_SOURCE = EVIDENCE_INPUT / "10_claim_evidence_index.csv"
REPORT_MANIFEST_SOURCE = EVIDENCE_INPUT / "REPORT_EVIDENCE_MANIFEST.json"
MAIN_ALGORITHM_SOURCE = EVIDENCE_INPUT / "MAIN_ALGORITHM.json"
ALGORITHM_REGISTRY_SOURCE = EVIDENCE_INPUT / "ALGORITHM_REGISTRY.csv"
RAW_MANIFEST_SOURCE = EVIDENCE_INPUT / "RAW_MANIFEST.json"


def require(path: Path, kind: str = "file") -> None:
    exists = path.is_file() if kind == "file" else path.is_dir()
    if not exists:
        raise FileNotFoundError(f"missing required {kind}: {path}")


require(BASELINE)
require(MML2OMML)

MATHML_NS = "http://www.w3.org/1998/Math/MathML"
_MML_TRANSFORM = etree.XSLT(etree.parse(str(MML2OMML)))


def file_sha256(path: Path) -> str:
    require(path)
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def read_parameter_design() -> list[dict[str, float | str]]:
    # 从正式配置中读取参数工况，不在绘图脚本中复制数值。
    require(PARAMETER_CONFIG)
    source = PARAMETER_CONFIG.read_text(encoding="utf-8")
    tree = ast.parse(source)
    scales: dict[str, tuple[float, float, float]] = {}
    for node in tree.body:
        if not isinstance(node, ast.Assign):
            continue
        if not any(isinstance(target, ast.Name) and target.id == "SCENE05B_SCALES" for target in node.targets):
            continue
        if not isinstance(node.value, ast.Dict):
            raise ValueError("SCENE05B_SCALES is not a dictionary")
        for key_node, value_node in zip(node.value.keys, node.value.values):
            case_id = ast.literal_eval(key_node)
            if isinstance(value_node, ast.Tuple):
                scales[case_id] = tuple(float(value) for value in ast.literal_eval(value_node))
            elif case_id == "Scene05B_Nominal":
                scales[case_id] = (1.0, 1.0, 1.0)
            else:
                raise ValueError(f"unsupported parameter expression: {case_id}")
        break
    expected_loop = 'for _bits in ("000", "001", "010", "011", "100", "101", "110", "111"):'
    if expected_loop not in source:
        raise ValueError("parameter corner construction changed in frozen runner")
    for bits in ("000", "001", "010", "011", "100", "101", "110", "111"):
        scales[f"Scene05B_C{bits}"] = tuple(0.9 if bit == "0" else 1.1 for bit in bits)
    required_cases = {
        "Scene05B_Nominal",
        "Scene05B_LiftMinus10",
        "Scene05B_PayloadPlus10",
        *(f"Scene05B_C{bits}" for bits in ("000", "001", "010", "011", "100", "101", "110", "111")),
    }
    if set(scales) != required_cases:
        raise ValueError(f"unexpected parameter cases: {sorted(scales)}")
    rows: list[dict[str, float | str]] = []
    for case_id, (lift_scale, mass_scale, inertia_scale) in scales.items():
        if case_id == "Scene05B_Nominal":
            coverage_class = "nominal"
        elif case_id in {"Scene05B_LiftMinus10", "Scene05B_PayloadPlus10"}:
            coverage_class = "official_focus"
        else:
            coverage_class = "corner"
        rows.append(
            {
                "case_id": case_id,
                "coverage_class": coverage_class,
                "lift_scale": lift_scale,
                "mass_scale": mass_scale,
                "inertia_scale": inertia_scale,
            }
        )
    return rows


SCENE_FIELDS = [
    ("standard_steps", "三轴标准阶跃"),
    ("scene01_piecewise_climb", "分段爬升"),
    ("scene02_spiral_climb", "螺旋爬升"),
    ("scene03_figure8", "八字轨迹"),
    ("scene04_hover", "定点悬停"),
    ("scene06b_disturbance", "三事件外力扰动"),
    ("parameter_11", "11项参数变化"),
    ("wind_5", "五相位风场"),
    ("sensor_5", "五工况传感器退化"),
    ("formation_3", "三机编队"),
]


def read_coverage_rows() -> list[dict[str, str]]:
    require(COVERAGE_SOURCE)
    with COVERAGE_SOURCE.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    expected = ["PID", "RA-GCA/Base", "RA-GCA-CGHTE", "CAP-ADRC", "CP-INDI"]
    if [row["controller"] for row in rows] != expected:
        raise ValueError("five-controller identity or order changed")
    for row in rows:
        available = [label for field, label in SCENE_FIELDS if row[field] != "not_available"]
        unavailable = [label for field, label in SCENE_FIELDS if row[field] == "not_available"]
        row["available_zh"] = "、".join(available) if available else "无"
        row["unavailable_zh"] = "、".join(unavailable) if unavailable else "无"
    return rows


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_supporting_tables() -> dict[str, Path]:
    # 表格底稿统一写为CSV，供Word和后续审计共同使用。
    required = [
        COVERAGE_SOURCE,
        STEP_SOURCE,
        COMMON_SOURCE,
        EVOLUTION_SOURCE,
        PARAMETER_SOURCE,
        REGRESSION_SOURCE,
        FORMATION_SOURCE,
        CLAIM_SOURCE,
        REPORT_MANIFEST_SOURCE,
        MAIN_ALGORITHM_SOURCE,
        ALGORITHM_REGISTRY_SOURCE,
        RAW_MANIFEST_SOURCE,
        PACKAGE_MANIFEST,
    ]
    for path in required:
        require(path)
    EVIDENCE_OUTPUT.mkdir(parents=True, exist_ok=True)

    coverage_rows = read_coverage_rows()
    coverage_output = EVIDENCE_OUTPUT / "chapter5_controller_coverage.csv"
    coverage_fields = [
        "controller",
        "role_zh",
        *[field for field, _ in SCENE_FIELDS],
        "available_zh",
        "unavailable_zh",
        "evidence_boundary_zh",
    ]
    write_csv(coverage_output, coverage_fields, coverage_rows)

    protocol_rows = [
        {
            "protocol_id": "P1",
            "comparison_zh": "赛题基线比较",
            "objects": "PID;RA-GCA-CGHTE",
            "scenes_zh": "三轴标准阶跃；分段爬升；螺旋爬升；八字轨迹；定点悬停；三事件外力扰动",
            "evidence_files": "02_standard_step_pid_vs_main.csv;03_five_controller_common_scenes.csv",
            "note_zh": "比较赛题基线与本文方案",
        },
        {
            "protocol_id": "P2",
            "comparison_zh": "五种方法比较",
            "objects": "PID;RA-GCA/Base;RA-GCA-CGHTE;CAP-ADRC;CP-INDI",
            "scenes_zh": "分段爬升；八字轨迹；定点悬停；三事件外力扰动",
            "evidence_files": "03_five_controller_common_scenes.csv",
            "note_zh": "比较不同控制思路",
        },
        {
            "protocol_id": "P3",
            "comparison_zh": "几何控制增强对比",
            "objects": "RA-GCA/Base;RA-GCA-CGHTE",
            "scenes_zh": "分段爬升；八字轨迹；定点悬停；三事件外力扰动",
            "evidence_files": "03_five_controller_common_scenes.csv",
            "note_zh": "比较基础方案与增强方案",
        },
        {
            "protocol_id": "P4",
            "comparison_zh": "编队安全比较",
            "objects": "相同RA-GCA-CGHTE低层控制器；PP-CBF关闭/开启",
            "scenes_zh": "三机近距离交汇",
            "evidence_files": "09_formation_pp_cbf.csv",
            "note_zh": "比较安全距离与跟踪表现",
        },
    ]
    protocol_output = EVIDENCE_OUTPUT / "chapter5_protocol_register.csv"
    write_csv(protocol_output, list(protocol_rows[0]), protocol_rows)

    main_algorithm = json.loads(MAIN_ALGORITHM_SOURCE.read_text(encoding="utf-8-sig"))
    key_rows = [
        {"item": "正文低层对象", "value": 5, "unit": "个", "definition": "PID、RA-GCA/Base、RA-GCA-CGHTE、CAP-ADRC、CP-INDI", "source": COVERAGE_SOURCE.name},
        {"item": "五对象公共场景", "value": 4, "unit": "个", "definition": "五个对象均有统一口径数据的公共场景", "source": COMMON_SOURCE.name},
        {"item": "控制周期", "value": 0.01, "unit": "s", "definition": "统一仿真输出与控制周期", "source": MAIN_ALGORITHM_SOURCE.name},
        {"item": "接口维度", "value": "11/18/2/4/16", "unit": "维", "definition": "参考/状态/分配限制/电机命令/诊断", "source": MAIN_ALGORITHM_SOURCE.name},
        {"item": "激活协方差阈值", "value": main_algorithm["fixed_parameters"]["activation_covariance_max"], "unit": "-", "definition": "正式发布参数", "source": MAIN_ALGORITHM_SOURCE.name},
        {"item": "激活持续时间", "value": main_algorithm["fixed_parameters"]["activation_persistence_s"], "unit": "s", "definition": "正式发布参数", "source": MAIN_ALGORITHM_SOURCE.name},
        {"item": "尺度应用速率", "value": main_algorithm["fixed_parameters"]["scale_rate_per_s"], "unit": "s^-1", "definition": "正式发布参数", "source": MAIN_ALGORITHM_SOURCE.name},
        {"item": "每步尺度变化上限", "value": main_algorithm["fixed_parameters"]["scale_rate_per_step"], "unit": "-", "definition": "正式发布参数", "source": MAIN_ALGORITHM_SOURCE.name},
    ]
    key_output = EVIDENCE_OUTPUT / "chapter5_key_numbers.csv"
    write_csv(key_output, list(key_rows[0]), key_rows)

    manifest = json.loads(PACKAGE_MANIFEST.read_text(encoding="utf-8-sig"))
    package_hashes = {item["path"]: item["sha256"] for item in manifest["files"]}
    official_model = "01_models/official/QuadrotorModel/package.mo"
    controller_paths = {
        "P1": [official_model, "01_models/sysblock/A8FormalRAGCACGHTE_20260715.mo"],
        "P2": [
            official_model,
            "01_models/comparators/ra_gca_base/A8FormalSO3V1F_20260711.mo",
            "01_models/sysblock/A8FormalRAGCACGHTE_20260715.mo",
            "01_models/comparators/cap_adrc/A8FormalADRCV11PX4AllocV1A_20260712.mo",
            "01_models/comparators/cp_indi/A8FormalINDICommandProxyV1A_20260712.mo",
        ],
        "P3": [
            "01_models/comparators/ra_gca_base/A8FormalSO3V1F_20260711.mo",
            "01_models/sysblock/A8FormalRAGCACGHTE_20260715.mo",
        ],
        "P4": [
            "01_models/sysblock/A8FormalRAGCACGHTE_20260715.mo",
            "01_models/sysblock/formation/A8FormalFormationPredictiveCBFV5C_20260712.mo",
        ],
    }
    lineage_rows = []
    for protocol in protocol_rows:
        paths = controller_paths[protocol["protocol_id"]]
        missing = [path for path in paths if path not in package_hashes]
        if missing:
            raise ValueError(f"controller paths missing from package manifest: {missing}")
        evidence_names = protocol["evidence_files"].split(";")
        evidence_paths = [EVIDENCE_INPUT / name for name in evidence_names]
        for evidence_path in evidence_paths:
            require(evidence_path)
        lineage_rows.append(
            {
                "protocol_id": protocol["protocol_id"],
                "algorithms": protocol["objects"],
                "scenes": protocol["scenes_zh"],
                "model_or_controller_files": ";".join(paths),
                "model_or_controller_sha256": ";".join(package_hashes[path] for path in paths),
                "solver": "DASSL",
                "sample_period_s": "0.01",
                "evaluation_window": "各场景采用相同评价区间",
                "raw_or_snapshot_evidence": protocol["evidence_files"],
                "metric_script": "02_scripts/evaluation/build_report_evidence.py;02_scripts/syslab/build_five_scene_comparison.jl",
                "source_evidence_sha256": ";".join(file_sha256(path) for path in evidence_paths),
                "output_hash_manifest": "chapter5_supporting_outputs_sha256.csv",
            }
        )
    lineage_output = EVIDENCE_OUTPUT / "chapter5_data_lineage.csv"
    write_csv(lineage_output, list(lineage_rows[0]), lineage_rows)

    outputs = {
        "coverage": coverage_output,
        "protocol": protocol_output,
        "key_numbers": key_output,
        "lineage": lineage_output,
    }
    hash_output = EVIDENCE_OUTPUT / "chapter5_supporting_outputs_sha256.csv"
    hash_rows = [
        {"artifact": path.name, "sha256": file_sha256(path)}
        for path in outputs.values()
    ]
    write_csv(hash_output, ["artifact", "sha256"], hash_rows)
    outputs["hashes"] = hash_output
    return outputs


def configure_plot_font() -> str:
    preferred = ["SimSun", "Microsoft YaHei", "Noto Sans CJK SC"]
    installed = {item.name for item in font_manager.fontManager.ttflist}
    selected = next((name for name in preferred if name in installed), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected]
    plt.rcParams["axes.unicode_minus"] = False
    return selected


def read_numeric_columns(path: Path, columns: tuple[str, ...], max_points: int = 1600) -> dict[str, list[float]]:
    # 长序列等距抽样并保留末点，避免改变终态信息。
    require(path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError(f"CSV has no header: {path}")
        missing = [column for column in columns if column not in reader.fieldnames]
        if missing:
            raise ValueError(f"missing columns {missing} in {path}")
        rows = list(reader)

    stride = max(1, math.ceil(len(rows) / max_points))
    sampled = rows[::stride]
    if rows and sampled[-1] is not rows[-1]:
        sampled.append(rows[-1])
    return {
        column: [float(row[column]) for row in sampled]
        for column in columns
    }


def style_plot_axis(ax) -> None:
    ax.grid(True, linestyle=":", linewidth=0.6, color="#AAB2B8", alpha=0.7)
    ax.tick_params(labelsize=8.2, colors="#30373D")
    for side in ("top", "right"):
        if side in ax.spines:
            ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        if side in ax.spines:
            ax.spines[side].set_color("#6D777E")
            ax.spines[side].set_linewidth(0.7)


def add_panel_label(ax, label: str) -> None:
    text_method = ax.text2D if hasattr(ax, "text2D") else ax.text
    text_method(
        0.01,
        0.98,
        label,
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=9.2,
        fontweight="bold",
        color="#1F252B",
    )


def save_figure_set(fig, stem: str) -> tuple[Path, Path, Path]:
    BUILD.mkdir(parents=True, exist_ok=True)
    png = BUILD / f"{stem}.png"
    pdf = BUILD / f"{stem}.pdf"
    svg = BUILD / f"{stem}.svg"
    fig.savefig(png, dpi=450, bbox_inches="tight", pad_inches=0.04, facecolor="white")
    fig.savefig(pdf, bbox_inches="tight", pad_inches=0.04, facecolor="white")
    fig.savefig(svg, bbox_inches="tight", pad_inches=0.04, facecolor="white")
    plt.close(fig)
    return png, pdf, svg


def generate_analysis_flow_figure() -> tuple[Path, Path, Path]:
    BUILD.mkdir(parents=True, exist_ok=True)
    configure_plot_font()
    png = BUILD / "figure_5_3_analysis_flow.png"
    pdf = BUILD / "figure_5_3_analysis_flow.pdf"
    svg = BUILD / "figure_5_3_analysis_flow.svg"

    fig, ax = plt.subplots(figsize=(8.0, 2.25))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2.7)
    ax.axis("off")

    labels = [
        ("实验设置", "模型、初值与参考\n求解器和采样周期"),
        ("Sysplorer仿真", "运行场景\n导出时序数据"),
        ("数据检查", "检查字段与单位\n确定评价区间"),
        ("Syslab分析", "赛题指标\n补充评价指标"),
        ("结果展示", "绘制曲线与表格\n分析实验结果"),
    ]
    colors = ["#FFFFFF", "#F5F5F5", "#ECECEC", "#F5F5F5", "#FFFFFF"]
    edge_colors = ["#4A4A4A"] * len(labels)
    x_positions = [0.22, 2.22, 4.22, 6.22, 8.22]
    bounded_text = []
    for index, ((title, subtitle), x, fill, edge) in enumerate(zip(labels, x_positions, colors, edge_colors)):
        box = FancyBboxPatch(
            (x, 1.18),
            1.56,
            0.92,
            boxstyle="round,pad=0.018,rounding_size=0.035",
            linewidth=0.9,
            edgecolor=edge,
            facecolor=fill,
        )
        ax.add_patch(box)
        title_text = ax.text(
            x + 0.78,
            1.78,
            title,
            ha="center",
            va="center",
            fontsize=9.6,
            fontweight="bold",
            color="#111111",
        )
        subtitle_text = ax.text(
            x + 0.78,
            1.46,
            subtitle,
            ha="center",
            va="center",
            fontsize=7.8,
            color="#333333",
            linespacing=1.18,
        )
        bounded_text.extend(((box, title_text), (box, subtitle_text)))
        if index < len(labels) - 1:
            arrow = FancyArrowPatch(
                (x + 1.60, 1.64),
                (x_positions[index + 1] - 0.05, 1.64),
                arrowstyle="-|>",
                mutation_scale=10,
                linewidth=0.9,
                color="#4A4A4A",
            )
            ax.add_patch(arrow)

    guard = FancyBboxPatch(
        (0.48, 0.34),
        9.04,
        0.42,
        boxstyle="round,pad=0.012,rounding_size=0.025",
        linewidth=0.7,
        edgecolor="#5A5A5A",
        facecolor="#FFFFFF",
    )
    ax.add_patch(guard)
    guard_text = ax.text(
        5.0,
        0.55,
        "所有方法采用相同的模型、初始条件、参考输入、求解器、采样周期和评价区间",
        ha="center",
        va="center",
        fontsize=8.8,
        color="#333333",
    )
    bounded_text.append((guard, guard_text))

    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    for patch, text_artist in bounded_text:
        patch_box = patch.get_window_extent(renderer)
        text_box = text_artist.get_window_extent(renderer)
        margin = 3.0
        assert text_box.x0 >= patch_box.x0 + margin
        assert text_box.x1 <= patch_box.x1 - margin
        assert text_box.y0 >= patch_box.y0 + margin
        assert text_box.y1 <= patch_box.y1 - margin

    fig.subplots_adjust(left=0.01, right=0.99, top=0.98, bottom=0.04)
    fig.savefig(png, dpi=450, bbox_inches="tight", pad_inches=0.03, facecolor="white")
    fig.savefig(pdf, bbox_inches="tight", pad_inches=0.03, facecolor="white")
    fig.savefig(svg, bbox_inches="tight", pad_inches=0.03, facecolor="white")
    plt.close(fig)
    return png, pdf, svg


def generate_reference_tasks_figure() -> tuple[Path, Path, Path]:
    configure_plot_font()
    step_x = read_numeric_columns(STEP_X, ("time", "referenceVector[1]"), 1000)
    step_y = read_numeric_columns(STEP_Y, ("time", "referenceVector[2]"), 1000)
    step_z = read_numeric_columns(STEP_Z, ("time", "referenceVector[3]"), 1000)
    climb = read_numeric_columns(SCENE_01, ("time", "referenceVector[3]"), 1200)
    spiral = read_numeric_columns(
        SCENE_02,
        ("referenceVector[1]", "referenceVector[2]", "referenceVector[3]"),
        1800,
    )
    figure8 = read_numeric_columns(SCENE_03, ("referenceVector[1]", "referenceVector[2]"), 1800)

    fig = plt.figure(figsize=(8.0, 5.5), facecolor="white")
    grid = fig.add_gridspec(2, 2, wspace=0.27, hspace=0.30)

    ax_step = fig.add_subplot(grid[0, 0])
    ax_step.plot(step_x["time"], step_x["referenceVector[1]"], color="#1F6F8B", linewidth=1.6, label="X轴参考")
    ax_step.plot(step_y["time"], step_y["referenceVector[2]"], color="#B7791F", linewidth=1.4, linestyle="--", label="Y轴参考")
    ax_step.plot(step_z["time"], step_z["referenceVector[3]"], color="#3D7A4C", linewidth=1.4, linestyle="-.", label="Z轴参考")
    ax_step.set_title("三轴独立标准阶跃", fontsize=10.2, pad=5)
    ax_step.set_xlabel("时间 / s", fontsize=8.8)
    ax_step.set_ylabel("参考位置 / m", fontsize=8.8)
    ax_step.legend(loc="best", frameon=False, fontsize=7.7, ncol=1)
    style_plot_axis(ax_step)
    add_panel_label(ax_step, "(a)")

    ax_climb = fig.add_subplot(grid[0, 1])
    ax_climb.plot(climb["time"], climb["referenceVector[3]"], color="#7A4EAB", linewidth=1.6)
    ax_climb.set_title("阶梯爬升高度参考", fontsize=10.2, pad=5)
    ax_climb.set_xlabel("时间 / s", fontsize=8.8)
    ax_climb.set_ylabel("高度参考 / m", fontsize=8.8)
    style_plot_axis(ax_climb)
    add_panel_label(ax_climb, "(b)")

    ax_spiral = fig.add_subplot(grid[1, 0], projection="3d")
    ax_spiral.plot(
        spiral["referenceVector[1]"],
        spiral["referenceVector[2]"],
        spiral["referenceVector[3]"],
        color="#1F6F8B",
        linewidth=1.5,
    )
    ax_spiral.scatter(
        [spiral["referenceVector[1]"][0], spiral["referenceVector[1]"][-1]],
        [spiral["referenceVector[2]"][0], spiral["referenceVector[2]"][-1]],
        [spiral["referenceVector[3]"][0], spiral["referenceVector[3]"][-1]],
        color=["#3D7A4C", "#B24C45"],
        s=22,
        depthshade=False,
    )
    ax_spiral.set_title("螺旋爬升空间参考", fontsize=10.2, pad=4)
    ax_spiral.set_xlabel("X / m", fontsize=8.2, labelpad=2)
    ax_spiral.set_ylabel("Y / m", fontsize=8.2, labelpad=2)
    ax_spiral.set_zlabel("Z / m", fontsize=8.2, labelpad=2)
    ax_spiral.tick_params(labelsize=7.5, pad=0)
    ax_spiral.view_init(elev=24, azim=-58)
    ax_spiral.set_box_aspect((1.0, 1.0, 0.72))
    add_panel_label(ax_spiral, "(c)")

    ax_figure8 = fig.add_subplot(grid[1, 1])
    ax_figure8.plot(figure8["referenceVector[1]"], figure8["referenceVector[2]"], color="#B7791F", linewidth=1.6)
    ax_figure8.scatter(
        [figure8["referenceVector[1]"][0], figure8["referenceVector[1]"][-1]],
        [figure8["referenceVector[2]"][0], figure8["referenceVector[2]"][-1]],
        color=["#3D7A4C", "#B24C45"],
        s=24,
        zorder=3,
    )
    ax_figure8.set_title("8字平面参考", fontsize=10.2, pad=5)
    ax_figure8.set_xlabel("X / m", fontsize=8.8)
    ax_figure8.set_ylabel("Y / m", fontsize=8.8)
    ax_figure8.set_aspect("equal", adjustable="datalim")
    style_plot_axis(ax_figure8)
    add_panel_label(ax_figure8, "(d)")

    fig.subplots_adjust(left=0.08, right=0.98, top=0.96, bottom=0.08)
    return save_figure_set(fig, "figure_5_1_reference_tasks")


def contiguous_active_intervals(time: list[float], magnitude: list[float]) -> list[tuple[float, float]]:
    intervals: list[tuple[float, float]] = []
    start: float | None = None
    for current_time, value in zip(time, magnitude):
        active = value > 1e-9
        if active and start is None:
            start = current_time
        elif not active and start is not None:
            intervals.append((start, current_time))
            start = None
    if start is not None and time:
        intervals.append((start, time[-1]))
    return intervals


def generate_stress_inputs_figure() -> tuple[Path, Path, Path]:
    configure_plot_font()
    parameter_rows = read_parameter_design()
    disturbance = read_numeric_columns(
        SCENE_06B,
        ("time", "quadChassisTest17_1.perRotorDisturbance"),
        2400,
    )

    fig = plt.figure(figsize=(8.0, 3.75), facecolor="white")
    grid = fig.add_gridspec(1, 2, width_ratios=(0.95, 1.35), wspace=0.36)

    ax_param = fig.add_subplot(grid[0, 0], projection="3d")
    limits = (0.9, 1.1)
    vertices = [
        (mass, lift, inertia)
        for mass in limits
        for lift in limits
        for inertia in limits
    ]
    for index, first in enumerate(vertices):
        for second in vertices[index + 1:]:
            changed = sum(abs(a - b) > 1e-12 for a, b in zip(first, second))
            if changed == 1:
                ax_param.plot(
                    [first[0], second[0]],
                    [first[1], second[1]],
                    [first[2], second[2]],
                    color="#AAB2B8",
                    linewidth=0.8,
                    linestyle=":",
                )

    corners = [row for row in parameter_rows if row["coverage_class"] == "corner"]
    focus = [row for row in parameter_rows if row["coverage_class"] == "official_focus"]
    ax_param.scatter(
        [float(row["mass_scale"]) for row in corners],
        [float(row["lift_scale"]) for row in corners],
        [float(row["inertia_scale"]) for row in corners],
        s=30,
        color="#5B6770",
        marker="o",
        depthshade=False,
        label="八角点组合",
    )
    ax_param.scatter(
        [float(row["mass_scale"]) for row in focus],
        [float(row["lift_scale"]) for row in focus],
        [float(row["inertia_scale"]) for row in focus],
        s=42,
        color="#1F6F8B",
        marker="D",
        depthshade=False,
        label="重点工况",
    )
    ax_param.scatter([1.0], [1.0], [1.0], s=58, color="#D59B20", marker="*", depthshade=False, label="名义参数")
    ax_param.set_title("参数变化设计点", fontsize=10.2, pad=5)
    ax_param.set_xlabel("质量系数", fontsize=8.2, labelpad=3)
    ax_param.set_ylabel("升力系数", fontsize=8.2, labelpad=3)
    ax_param.set_zlabel("")
    ax_param.text2D(
        1.050,
        0.50,
        "惯量系数",
        transform=ax_param.transAxes,
        rotation=90,
        ha="center",
        va="center",
        fontsize=8.2,
        color="#222222",
    )
    ax_param.set_xticks([0.9, 1.0, 1.1])
    ax_param.set_yticks([0.9, 1.0, 1.1])
    ax_param.set_zticks([0.9, 1.0, 1.1])
    ax_param.tick_params(labelsize=7.4, pad=0)
    ax_param.view_init(elev=23, azim=-52)
    ax_param.set_box_aspect((1, 1, 0.82))
    ax_param.legend(loc="upper left", bbox_to_anchor=(-0.05, 0.96), frameon=False, fontsize=7.2)
    add_panel_label(ax_param, "(a)")

    ax_force = fig.add_subplot(grid[0, 1])
    time = disturbance["time"]
    per_rotor_force = disturbance["quadChassisTest17_1.perRotorDisturbance"]
    ax_force.step(time, per_rotor_force, where="post", color="#1F6F8B", linewidth=1.45)
    intervals = contiguous_active_intervals(time, per_rotor_force)
    upper = max(max(per_rotor_force), 1e-6)
    for event_index, (start, end) in enumerate(intervals, start=1):
        ax_force.axvspan(start, end, color="#D8DEE2", alpha=0.30, linewidth=0)
        ax_force.text(
            (start + end) / 2,
            upper * 0.90,
            f"事件{event_index}",
            ha="center",
            va="top",
            fontsize=7.4,
            color="#4B555C",
        )
    ax_force.set_title("三事件外力输入", fontsize=10.2, pad=5)
    ax_force.set_xlabel("时间 / s", fontsize=8.8)
    ax_force.set_ylabel("X向力 / N", fontsize=8.8)
    ax_force.set_ylim(-0.0015, upper * 1.14)
    style_plot_axis(ax_force)
    add_panel_label(ax_force, "(b)")

    fig.subplots_adjust(left=0.04, right=0.985, top=0.95, bottom=0.11)
    return save_figure_set(fig, "figure_5_2_stress_inputs")


def set_style_fonts(style, east_asia: str = "宋体", latin: str = "Times New Roman") -> None:
    style.font.name = latin
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("w:ascii", latin), ("w:hAnsi", latin),
                       ("w:eastAsia", east_asia), ("w:cs", latin)):
        rfonts.set(qn(key), value)


def set_run_fonts(run, size: float | None = None, bold: bool | None = None,
                  east_asia: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("w:ascii", latin), ("w:hAnsi", latin),
                       ("w:eastAsia", east_asia), ("w:cs", latin)):
        rfonts.set(qn(key), value)


def clear_body_keep_section(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._element):
        paragraph._element.remove(child)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style_fonts(normal)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 15, 0, 12),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 10, 4),
    ):
        style = doc.styles[name]
        set_style_fonts(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    set_style_fonts(caption)
    caption.font.size = Pt(10.5)
    caption.font.bold = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_with_next = True

    for style_name, size, left, first in (
        ("Source Note", 9, 0, 0),
        ("Reference Entry", 10.5, 24, -24),
    ):
        try:
            style = doc.styles[style_name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_fonts(style)
        style.font.size = Pt(size)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.left_indent = Pt(left)
        style.paragraph_format.first_line_indent = Pt(first)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

    header = section.header
    hp = header.paragraphs[0]
    clear_paragraph(hp)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("A8 四旋翼无人机位姿控制系统设计优化  仿真分析报告第5章")
    set_run_fonts(hr, 9)
    ppr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    borders.append(bottom)
    ppr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr1 = fp.add_run("第 ")
    set_run_fonts(fr1, 9)
    add_field(fp, "PAGE")
    fr2 = fp.add_run(" 页")
    set_run_fonts(fr2, 9)


def add_body(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    run = paragraph.add_run(text)
    set_run_fonts(run, 12)
    return paragraph


def add_body_segments(doc: Document, segments):
    paragraph = doc.add_paragraph(style="Normal")
    for segment in segments:
        text, italic = segment[:2]
        subscript = segment[2] if len(segment) > 2 else False
        run = paragraph.add_run(text)
        set_run_fonts(run, 12)
        run.italic = italic
        run.font.subscript = subscript
    return paragraph


def add_heading(doc: Document, text: str, level: int, page_break: bool = False):
    paragraph = doc.add_heading(text, level=level)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Caption")
    run = paragraph.add_run(text)
    set_run_fonts(run, 10.5, True)
    return paragraph


def add_source_note(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Source Note")
    run = paragraph.add_run(text)
    set_run_fonts(run, 9)
    return paragraph


def set_cell_text(cell, text: str, bold: bool = False,
                  align=WD_ALIGN_PARAGRAPH.LEFT, size: float = 9.2) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_fonts(run, size, bold)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def set_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, value, size in (
        ("top", "single", "8"),
        ("bottom", "single", "8"),
        ("left", "nil", "0"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)

    header = table.rows[0]
    for cell in header.cells:
        tcpr = cell._tc.get_or_add_tcPr()
        cell_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        cell_borders.append(bottom)
        tcpr.append(cell_borders)


def set_cell_margins(cell, top=50, start=65, bottom=50, end=65) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_three_line_table(doc: Document, caption: str, headers: list[str],
                         rows: list[list[str]], widths_cm: list[float],
                         font_size: float = 9.2,
                         page_break_before: bool = False):
    caption_paragraph = add_caption(doc, caption)
    if page_break_before:
        caption_paragraph.paragraph_format.page_break_before = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        table.columns[index].width = Cm(widths_cm[index])
        set_cell_text(table.cell(0, index), header, True, WD_ALIGN_PARAGRAPH.CENTER, font_size)
    for row_values in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(row_values):
            cells[index].width = Cm(widths_cm[index])
            set_cell_text(cells[index], value, False, WD_ALIGN_PARAGRAPH.LEFT, font_size)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
    set_table_borders(table)
    return table


def mathml_to_omml(mathml: str):
    source = etree.fromstring(mathml.encode("utf-8"))
    converted = _MML_TRANSFORM(source)
    return parse_xml(etree.tostring(converted.getroot()))


def add_native_equation(doc: Document, mathml: str, number: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(13.3)
    table.columns[1].width = Cm(2.0)
    for cell, width in zip(table.rows[0].cells, (Cm(13.3), Cm(2.0))):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        tc_width = tcpr.find(qn("w:tcW"))
        if tc_width is None:
            tc_width = OxmlElement("w:tcW")
            tcpr.insert(0, tc_width)
        tc_width.set(qn("w:w"), str(width.twips))
        tc_width.set(qn("w:type"), "dxa")
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "nil")
            borders.append(node)
        tcpr.append(borders)

    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left.paragraph_format.first_line_indent = Pt(0)
    left.paragraph_format.space_before = Pt(3)
    left.paragraph_format.space_after = Pt(3)
    left._p.append(mathml_to_omml(mathml))

    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.first_line_indent = Pt(0)
    run = right.add_run(f"（{number}）")
    set_run_fonts(run, 11)


def add_figure(doc: Document, image_path: Path, caption: str, alt_text: str) -> None:
    require(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Cm(15.4))
    inline._inline.docPr.set("descr", alt_text)
    add_caption(doc, caption)


def add_reference(doc: Document, segments: list[tuple[str, bool]]) -> None:
    paragraph = doc.add_paragraph(style="Reference Entry")
    for text, italic in segments:
        run = paragraph.add_run(text)
        set_run_fonts(run, 10.5)
        run.italic = italic


def build_document() -> Path:
    # 文档重建使用固定模板，并拒绝覆盖输入草稿。
    if OUTPUT.exists():
        raise FileExistsError(f"output already exists: {OUTPUT}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(parents=True, exist_ok=True)
    write_supporting_tables()
    figure_png, _, _ = generate_analysis_flow_figure()
    reference_tasks_png, _, _ = generate_reference_tasks_figure()
    stress_inputs_png, _, _ = generate_stress_inputs_figure()
    doc = Document(BASELINE)
    clear_body_keep_section(doc)
    setup_document(doc)
    doc.core_properties.title = "A8仿真分析报告第5章：实验设计与评价方法"
    doc.core_properties.subject = "五种控制器的对比条件、实验场景、评价指标和数据处理方法"
    doc.core_properties.author = "A8项目组"
    doc.core_properties.keywords = "四旋翼；实验设计；评价指标；数据处理；可复现性"

    add_heading(doc, "5  实验设计与评价方法", 1)
    add_body(
        doc,
        "本章说明后续性能比较采用的实验设置与评价方法。为保证结果可比，各算法使用相同的四旋翼模型、参考输入、初始条件、求解设置和评价区间，并按统一方法计算评价指标。第6章的单机性能比较和第7章的编队安全验证均遵循这些设置。",
    )

    add_heading(doc, "5.1  对比条件与场景设计", 2)
    add_body(
        doc,
        "单机实验比较五种方法。PID是赛题给出的基线，RA-GCA/Base是几何控制基础方案，RA-GCA-CGHTE为本文方案；CAP-ADRC和CP-INDI分别代表主动扰动抑制和增量动态逆方法。",
    )
    add_body(
        doc,
        "编队实验在三套RA-GCA-CGHTE控制器之上加入PP-CBF，用于比较安全监督开启前后的距离、跟踪误差和控制修正。PP-CBF不参与五种单机方法比较。",
    )
    add_body(
        doc,
        "RA-GCA-CGHTE实验采用第3章给出的控制参数，控制周期为0.01 s。",
    )
    add_body(
        doc,
        "实验分为四组：PID与RA-GCA-CGHTE比较赛题基线和本文方案；五种方法比较不同控制思路；RA-GCA/Base与RA-GCA-CGHTE比较几何控制增强效果；PP-CBF关闭与开启比较编队安全监督作用。",
    )
    add_body(
        doc,
        "同组实验使用相同的赛题四旋翼模型、物理参数、参考输入、初始状态、仿真时长、DASSL求解器、0.01 s采样周期、评价区间和指标定义。",
    )
    add_body(
        doc,
        "表5-1概括五种单机控制方法的基本思路和比较作用。",
    )
    comparison_rows = [
        ["PID", "级联比例-积分-微分控制", "赛题基线"],
        ["RA-GCA/Base", "约化姿态几何控制", "几何控制基础方案"],
        ["RA-GCA-CGHTE", "几何控制、推力估计与约束感知分配", "本文方案"],
        ["CAP-ADRC", "主动扰动抑制", "抗扰方法对照"],
        ["CP-INDI", "增量非线性动态逆", "动态逆方法对照"],
    ]
    add_three_line_table(
        doc,
        "表5-1  五种单机控制方法及比较作用",
        ["方法", "控制思路", "比较作用"],
        comparison_rows,
        [3.2, 7.4, 4.8],
        9.0,
    )

    add_body(
        doc,
        "表5-2列出四类比较及其使用的场景，分别展示相对PID的改进、不同控制思路、几何控制增强和编队安全监督。",
    )
    protocol_rows = [
        ["赛题基线比较", "PID、RA-GCA-CGHTE", "三轴阶跃、典型轨迹、定点悬停和外力扰动"],
        ["五种方法比较", "PID、RA-GCA/Base、RA-GCA-CGHTE、CAP-ADRC、CP-INDI", "分段爬升、8字轨迹、定点悬停和三事件外扰"],
        ["几何控制增强", "RA-GCA/Base、RA-GCA-CGHTE", "分段爬升、8字轨迹、定点悬停和三事件外扰"],
        ["编队安全比较", "RA-GCA-CGHTE、PP-CBF关闭/开启", "三机近距离交汇"],
    ]
    add_three_line_table(
        doc,
        "表5-2  四类比较安排与适用范围",
        ["比较内容", "方法", "场景"],
        protocol_rows,
        [3.3, 6.1, 6.0],
        8.8,
        page_break_before=True,
    )

    add_body(
        doc,
        "其余实验包括螺旋爬升、五相位风场、五工况传感器退化、三机编队保持与换位，用于分别考察三维轨迹、环境扰动、观测误差和协同飞行。跨方法比较只使用实验条件相同的结果。",
    )

    add_body(
        doc,
        "图5-1集中展示三轴独立标准阶跃、阶梯爬升、螺旋爬升和8字轨迹的参考输入。四类任务由单轴瞬态逐步扩展到多事件升降和三维耦合轨迹，用于观察不同控制器在响应速度、连续跟踪、转弯爬升和方向反转条件下的表现。",
    )
    add_figure(
        doc,
        reference_tasks_png,
        "图5-1  代表性任务的参考输入",
        "图5-1代表性任务参考：三轴独立标准阶跃、阶梯爬升、螺旋爬升和8字轨迹",
    )
    add_source_note(doc, "资料来源：根据各实验场景的参考输入整理。")
    add_body(
        doc,
        "标准阶跃便于分离各轴的瞬态特征；阶梯爬升增加多个高度变化事件；螺旋和8字轨迹分别引入三维耦合、重复转弯和方向反转。统一参考输入使不同控制器能够在相同任务条件下比较。",
    )
    add_body(
        doc,
        "图5-2给出参数变化设计点和三次外力事件。参数场景同时覆盖质量、升力系数和惯量变化；外力场景在四个扰动作用点同步施加X向力，右图只绘制单个作用点的输入，避免重复显示四条完全相同的曲线。",
    )
    add_figure(
        doc,
        stress_inputs_png,
        "图5-2  参数变化设计点与三事件单旋翼外力输入",
        "图5-2参数变化和外力输入：八角点组合、重点工况、名义参数以及三次单旋翼X向外力事件",
    )
    add_source_note(doc, "资料来源：根据正式场景配置和三事件外力仿真数据整理。")
    add_body(
        doc,
        "参数图由八个参数组合、两个重点工况和名义参数构成，可区分单一因素与组合变化；外力输入包含三个持续时间和幅值不同的事件。两幅图均描述预先设定的实验条件，不包含控制器响应或性能结论。",
    )
    add_body(
        doc,
        "同一场景中的各方法采用相同模型、参考输入和评价区间。",
    )

    add_heading(doc, "5.2  赛题要求指标", 2)
    add_body_segments(doc, [
        ("赛题要求使用超调量、调节时间、稳态误差和均方根误差（Root Mean Square Error, RMSE）定量评价优化前后的控制性能。为保证计算方法一致，本文采用以下离散定义。设阶跃时刻为 ", False),
        ("t₀", True),
        ("，阶跃幅值为 ", False),
        ("Δr", True),
        ("，最终目标为 ", False),
        ("r", True),
        ("f", True, True),
        ("，方向符号 ", False),
        ("s = sign(Δr)", True),
        ("，响应序列为 ", False),
        ("yₖ", True),
        ("，正部算子为 ", False),
        ("[x]₊ = max(x, 0)", True),
        ("。百分比超调量定义为", False),
    ])
    add_native_equation(
        doc,
        f"""<math xmlns="{MATHML_NS}" display="block"><mrow>
        <msub><mi>M</mi><mi>p</mi></msub><mo>=</mo><mn>100</mn><mo>·</mo>
        <mfrac><mrow><munder><mo>max</mo><mrow><mi>k</mi><mo>≥</mo><msub><mi>k</mi><mn>0</mn></msub></mrow></munder>
        <msub><mrow><mo>[</mo><mi>s</mi><mo>(</mo><msub><mi>y</mi><mi>k</mi></msub><mo>-</mo><msub><mi>r</mi><mi>f</mi></msub><mo>)</mo><mo>]</mo></mrow><mo>+</mo></msub></mrow>
        <mrow><mo>|</mo><mi>Δr</mi><mo>|</mo></mrow></mfrac></mrow></math>""",
        "5-1",
    )
    add_body_segments(doc, [
        ("三轴独立标准阶跃在10 s触发，幅值为1 m。调节带取 ", False),
        ("δ = max(0.02|Δr|, 0.02 m)", True),
        ("，响应进入调节带后需连续保持 ", False),
        ("T", True),
        ("h", True, True),
        (" = 2 s", True),
        ("。调节时间取满足持续保持条件的最早时刻与阶跃时刻之差：", False),
    ])
    add_native_equation(
        doc,
        f"""<math xmlns="{MATHML_NS}" display="block"><mrow>
        <msub><mi>t</mi><mi>s</mi></msub><mo>=</mo><mi>min</mi><mo>{{</mo>
        <msub><mi>t</mi><mi>k</mi></msub><mo>-</mo><msub><mi>t</mi><mn>0</mn></msub><mo>:</mo>
        <mo>|</mo><msub><mi>y</mi><mi>j</mi></msub><mo>-</mo><msub><mi>r</mi><mi>f</mi></msub><mo>|</mo><mo>≤</mo><mi>δ</mi><mo>,</mo>
        <mo>∀</mo><msub><mi>t</mi><mi>j</mi></msub><mo>∈</mo><mo>[</mo><msub><mi>t</mi><mi>k</mi></msub><mo>,</mo>
        <msub><mi>t</mi><mi>k</mi></msub><mo>+</mo><msub><mi>T</mi><mi>h</mi></msub><mo>]</mo><mo>}}</mo></mrow></math>""",
        "5-2",
    )
    add_body(
        doc,
        "稳态误差使用阶跃后末段20%样本，既保留有符号均值用于判断偏差方向，也计算绝对均值用于算法比较。正文主要报告稳态绝对误差：",
    )
    add_native_equation(
        doc,
        f"""<math xmlns="{MATHML_NS}" display="block"><mrow>
        <msub><mi>e</mi><mrow><mi>s</mi><mi>s</mi></mrow></msub><mo>=</mo>
        <mfrac><mn>1</mn><msub><mi>N</mi><mi>s</mi></msub></mfrac>
        <munder><mo>∑</mo><mrow><mi>k</mi><mo>∈</mo><msub><mi>W</mi><mi>s</mi></msub></mrow></munder>
        <mo>|</mo><msub><mi>r</mi><mi>f</mi></msub><mo>-</mo><msub><mi>y</mi><mi>k</mi></msub><mo>|</mo></mrow></math>""",
        "5-3",
    )
    add_body_segments(doc, [
        ("轨迹任务采用三维位置误差 ", False),
        ("eₖ = p", True),
        ("d,k", True, True),
        (" - pₖ", True),
        ("。全程RMSE同时反映三个方向的跟踪偏差，避免只选取表现较好的单轴：", False),
    ])
    add_native_equation(
        doc,
        f"""<math xmlns="{MATHML_NS}" display="block"><mrow><mi>RMSE</mi><mo>=</mo><msqrt><mrow>
        <mfrac><mn>1</mn><mi>N</mi></mfrac>
        <munderover><mo>∑</mo><mrow><mi>k</mi><mo>=</mo><mn>1</mn></mrow><mi>N</mi></munderover>
        <msup><mrow><mo>∥</mo><msub><mi>p</mi><mrow><mi>d</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>-</mo><msub><mi>p</mi><mi>k</mi></msub><mo>∥</mo></mrow><mn>2</mn></msup>
        </mrow></msqrt></mrow></math>""",
        "5-4",
    )
    add_body(
        doc,
        "式（5-1）至式（5-4）给出统一计算方法。2%调节带、0.02 m绝对带、2 s持续时间和末段20%数据为本文采用的计算约定，用于保证不同控制器之间的结果可比；这些数值并非赛题规定的阈值。",
    )

    add_heading(doc, "5.3  补充指标与判断条件", 2)
    add_body(
        doc,
        "四项赛题要求指标主要反映位置跟踪性能。复杂场景还需考察误差累积、恢复速度、电机指令变化、姿态状态和编队安全。表5-3按七类问题列出补充指标。控制分配残差和电机变化率反映电机指令是否平滑、期望控制量能否实现（Johansen & Fossen, 2013）；最小间距、违规时间和修正能量用于比较安全监督开启前后的安全性及任务影响（Ames et al., 2017）。",
    )
    metric_rows = [
        ["阶跃响应", "Mₚ / %；tₛ / s；eₛₛ / m"],
        ["轨迹跟踪", "RMSE、IAE、ISE、P95、峰值"],
        ["参数变化", "RMSE保持比ρ"],
        ["扰动恢复", "恢复时间 / s；末段误差 / m"],
        ["控制输入", "电机总变差、变化率、分配残差"],
        ["姿态与运动", "倾角、角速度、旋转矩阵质量"],
        ["编队与安全", "质心/队形RMSE、间距、违规时间、修正能量"],
    ]
    add_three_line_table(
        doc,
        "表5-3  赛题要求指标与补充指标",
        ["评价问题", "指标及单位"],
        metric_rows,
        [4.2, 11.6],
        8.8,
    )

    add_body(
        doc,
        "除四项赛题要求指标外，本文同时检查姿态变化和执行器限制，避免只依据位置误差判断控制效果。标准阶跃采用10%超调量和0.02 m稳态绝对误差作为性能参考；参数变化使用RMSE保持比、峰值误差和稳态误差；外力扰动报告恢复时间和末段误差；编队实验以0.60 m作为最小安全距离参考。这些数值用于统一本文的实验评价，并非赛题规定的阈值。",
    )
    add_body(
        doc,
        "五相位风场和五种传感器退化均采用预先设定的确定性工况。报告给出各工况指标，并使用均值、样本标准差和范围汇总差异。样本标准差只反映这些预设工况之间的变化，不代表随机总体的置信区间；因此，单次确定性仿真不进行显著性检验。",
    )

    add_heading(doc, "5.4  仿真数据处理方法", 2)
    add_body(
        doc,
        "图5-3概括仿真数据的处理过程。各方法先在相同的模型、初始条件、参考输入和求解设置下通过Sysplorer运行，导出参考轨迹、状态、执行器和诊断数据；随后使用Syslab检查字段与单位，确定评价区间，并按第5.2节和第5.3节的方法计算指标。",
    )
    add_figure(
        doc,
        figure_png,
        "图5-3  仿真数据处理与结果分析流程",
        "图5-3仿真数据处理与结果分析：实验设置、Sysplorer仿真、数据检查、Syslab分析和结果展示",
    )
    add_source_note(doc, "资料来源：根据本项目的Sysplorer仿真和Syslab数据分析过程整理。")
    add_body(
        doc,
        "仿真数据保留原始采样时序，并统一单位、评价区间和指标定义。确定性工况逐项统计；同一工况有多次运行时，再给出均值和离散程度。",
    )
    add_body(
        doc,
        "图表标明指标名称、单位和评价区间，曲线、表格和文字分析均使用同一组仿真结果。",
    )

    add_heading(doc, "参考文献", 1, page_break=True)
    add_reference(doc, [
        ("Ames, A. D., Xu, X., Grizzle, J. W., & Tabuada, P. (2017). Control barrier function based quadratic programs for safety critical systems. ", False),
        ("IEEE Transactions on Automatic Control, 62", True),
        ("(8), 3861-3876. https://doi.org/10.1109/TAC.2016.2638961", False),
    ])
    add_reference(doc, [
        ("Johansen, T. A., & Fossen, T. I. (2013). Control allocation - A survey. ", False),
        ("Automatica, 49", True),
        ("(5), 1087-1103. https://doi.org/10.1016/j.automatica.2013.01.035", False),
    ])

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
